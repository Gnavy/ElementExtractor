#!/usr/bin/env python3
"""
校验 Case1 已填信息收集表是否符合 template_row_catalog 规则。

用法：
  python tools/validate_collection_filled.py \\
    --catalog outputs/template_row_catalog.json \\
    --filled outputs/collection_filled.xlsx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path


def _norm(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _remark_has_citation(remark: str) -> bool:
    """备注须含源文件引用与原文摘录。"""
    e = remark or ""
    if "<引用>" in e and "</引用>" in e and "原文" in e:
        return True
    # 兼容宽松写法：来源/文件名 + 原文「」
    if ("原文" in e or "「" in e) and (
        ".docx" in e.lower()
        or ".pptx" in e.lower()
        or ".pdf" in e.lower()
        or "尽调" in e
        or "可研" in e
        or "http://" in e
        or "https://" in e
    ):
        return True
    return False


def validate(catalog_path: Path, filled_path: Path) -> dict:
    from openpyxl import load_workbook

    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    wb = load_workbook(filled_path, read_only=True, data_only=True)
    errors: list[str] = []
    warnings: list[str] = []
    stats = {"yes_no_rows": 0, "exclusive_groups": 0, "remarks_filled": 0}

    try:
        for sheet in catalog.get("sheets") or []:
            name = sheet.get("name") or "Sheet1"
            if name not in wb.sheetnames:
                errors.append(f"缺少 sheet: {name}")
                continue
            ws = wb[name]
            for section in sheet.get("sections") or []:
                for grp in section.get("indicator_groups") or []:
                    mode = grp.get("choice_mode") or "unknown"
                    rows = grp.get("rows") or []
                    ind = grp.get("indicator_name") or "?"

                    if mode == "yes_no":
                        for r in rows:
                            stats["yes_no_rows"] += 1
                            row_n = r["row"]
                            d = _norm(ws.cell(row_n, 4).value)
                            e = _norm(ws.cell(row_n, 5).value)
                            if d and d not in ("是", "否"):
                                errors.append(
                                    f"行{row_n}「{ind}」指标选择须为是/否，实际：{d!r}"
                                )
                            if d and not e:
                                warnings.append(f"行{row_n} 已填指标选择但备注为空")
                            elif d and e and not _remark_has_citation(e):
                                warnings.append(
                                    f"行{row_n}「{ind}」备注缺少源文件名与原文引用"
                                    f"（建议：判断要点。<引用>文件名：原文「…」</引用>）"
                                )
                            if d and e:
                                stats["remarks_filled"] += 1

                    elif mode == "exclusive":
                        stats["exclusive_groups"] += 1
                        filled_rows = []
                        for r in rows:
                            row_n = r["row"]
                            d = _norm(ws.cell(row_n, 4).value)
                            if d:
                                filled_rows.append((row_n, d, r.get("option_text", "")))
                        if len(filled_rows) == 0:
                            errors.append(f"互斥组「{ind}」未选择任何选项（应恰选 1 项）")
                        elif len(filled_rows) > 1:
                            errors.append(
                                f"互斥组「{ind}」有多行填了指标选择："
                                f"{[x[0] for x in filled_rows]}"
                            )
                        else:
                            row_n, d, opt = filled_rows[0]
                            e = _norm(ws.cell(row_n, 5).value)
                            allowed = grp.get("allowed_choices") or []
                            if allowed and d not in allowed and d != _norm(opt):
                                warnings.append(
                                    f"行{row_n} 指标选择 {d!r} 不在 allowed_choices 内"
                                )
                            if not e:
                                warnings.append(f"行{row_n} 已选选项但备注为空")
                            elif not _remark_has_citation(e):
                                warnings.append(
                                    f"行{row_n}「{ind}」备注缺少源文件名与原文引用"
                                    f"（建议：判断要点。<引用>文件名：原文「…」</引用>）"
                                )
                            if e:
                                stats["remarks_filled"] += 1
    finally:
        wb.close()

    passed = len(errors) == 0
    return {
        "passed": passed,
        "errors": errors,
        "warnings": warnings,
        "stats": stats,
    }


def main() -> None:
    p = argparse.ArgumentParser(description="校验 Case1 已填 xlsx")
    p.add_argument("--catalog", required=True)
    p.add_argument("--filled", required=True)
    p.add_argument("--out", default="outputs/validation_report.json")
    args = p.parse_args()

    cwd = Path.cwd().resolve()
    catalog = (cwd / args.catalog).resolve()
    filled = (cwd / args.filled).resolve()
    out = (cwd / args.out).resolve()
    for path in (catalog, filled):
        if not str(path).startswith(str(cwd)):
            raise SystemExit("路径必须位于当前工作目录之下")
    if not catalog.is_file() or not filled.is_file():
        raise SystemExit("catalog 或 filled 文件不存在")

    report = validate(catalog, filled)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"ok": report["passed"], **report}, ensure_ascii=False, indent=2))
    raise SystemExit(0 if report["passed"] else 1)


if __name__ == "__main__":
    main()
