#!/usr/bin/env python3
"""
列出 Excel 各 sheet 第 1 行表头，输出 JSON 到 stdout。
可选列出含公式的单元格坐标（填报收集表时需避让，勿覆盖）。
用法（ cwd 为解压项目根）：
  python tools/dump_sheet_headers.py --src inputs/collection_template.xlsx
  python tools/dump_sheet_headers.py --src inputs/collection_template.xlsx --formulas
依赖：openpyxl
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path


def _headers_row1(ws) -> list[str]:
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    cells = list(row) if row else []
    headers = []
    for c in cells:
        if c is None:
            headers.append("")
        else:
            headers.append(str(c).strip())
    return headers


def _formula_cells(ws) -> list[str]:
    refs: list[str] = []
    if not ws.max_row or not ws.max_column:
        return refs
    for row in ws.iter_rows(
        min_row=1,
        max_row=ws.max_row,
        min_col=1,
        max_col=ws.max_column,
    ):
        for cell in row:
            if cell.value is None and getattr(cell, "data_type", None) != "f":
                continue
            dt = getattr(cell, "data_type", None)
            if dt == "f":
                refs.append(cell.coordinate)
            elif isinstance(cell.value, str) and cell.value.startswith("="):
                refs.append(cell.coordinate)
    return sorted(set(refs))


def main() -> None:
    p = argparse.ArgumentParser(description="导出各 sheet 首行表头为 JSON")
    p.add_argument("--src", required=True, help="相对路径，如 inputs/collection_template.xlsx")
    p.add_argument(
        "--formulas",
        action="store_true",
        help="同时列出各 sheet 中含 Excel 公式的单元格坐标（勿手动覆盖）",
    )
    args = p.parse_args()

    cwd = Path.cwd().resolve()
    src = (cwd / args.src).resolve()
    if not str(src).startswith(str(cwd)):
        raise SystemExit("路径必须位于当前工作目录之下")
    if not src.is_file():
        raise SystemExit(f"文件不存在：{src}")

    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise SystemExit("需要 openpyxl：pip install openpyxl\n") from e

    sheets_out: dict = {}

    if args.formulas:
        wb = load_workbook(src, read_only=False, data_only=False)
        try:
            for name in wb.sheetnames:
                ws = wb[name]
                sheets_out[name] = {
                    "headers": _headers_row1(ws),
                    "formula_cells": _formula_cells(ws),
                }
        finally:
            wb.close()
    else:
        wb = load_workbook(src, read_only=True, data_only=True)
        try:
            for name in wb.sheetnames:
                ws = wb[name]
                sheets_out[name] = {"headers": _headers_row1(ws)}
        finally:
            wb.close()

    print(json.dumps({"sheets": sheets_out}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
