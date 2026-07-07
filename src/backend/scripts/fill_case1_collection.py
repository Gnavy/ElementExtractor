#!/usr/bin/env python3
"""
从 outputs/elements_extracted.json 应急回填 Case1 指标表（精度低于模板驱动填表）。
在智能体因 API 中断未写出 xlsx 时由 Worker 兜底调用。
若已存在 template_row_catalog.json，请续跑智能体按 SOP_信息收集表填报.md 填表，勿依赖本脚本。

用法（cwd = extract_root）：
  python tools/fill_case1_collection.py \\
    --template inputs/collection_template.xlsx \\
    --elements outputs/elements_extracted.json \\
    --out outputs/collection_filled.xlsx
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path


def _norm(s: str | None) -> str:
    if not s:
        return ""
    return re.sub(r"\s+", "", str(s).strip())


def _element_map(data: dict) -> dict[str, dict]:
    out: dict[str, dict] = {}
    for el in data.get("elements") or []:
        name = (el.get("element_name") or "").strip()
        if name:
            out[name] = el
    return out


def _remark_for_element(el: dict) -> str:
    logic = (el.get("logic_trace") or "").strip()
    situation = (el.get("situation") or "").strip()
    gaps = el.get("gaps") or []
    parts = []
    if logic:
        parts.append(f"口径依据：{logic}")
    if situation:
        parts.append(f"事实归纳：{situation[:800]}")
    if gaps:
        parts.append(f"缺口：{'；'.join(str(g) for g in gaps)}")
    parts.append("（本行指标选择由系统根据要素归纳做启发式匹配，请人工复核。）")
    return "；".join(parts)


def _option_core(text: str) -> str:
    t = text.strip()
    t = re.sub(r"^[□☑✓]\s*", "", t)
    return t


def _infer_yes_no(option_text: str, situation: str, logic_trace: str) -> str | None:
    """对 □ 开头的子选项行推断 是/否。"""
    blob = f"{situation} {logic_trace}"
    opt = _option_core(option_text)
    if not opt:
        return None
    neg_patterns = [
        r"未(见|有|落实|设立|约定|取得)",
        r"不存在",
        r"未能",
        r"无法",
        r"不包括",
        r"不算",
    ]
    for pat in neg_patterns:
        if re.search(pat, opt) and re.search(pat, blob):
            return "是"
    keywords = re.findall(r"[\u4e00-\u9fff]{2,}", opt)
    keywords = [k for k in keywords if len(k) >= 2][:12]
    hits = sum(1 for k in keywords if k in blob)
    if hits >= max(2, len(keywords) // 3):
        return "是"
    if any(w in blob for w in ("未", "无", "否", "未能", "无法")) and any(
        w in opt for w in ("未", "无", "不")
    ):
        return "否"
    return "否"


def _set_cell_value(ws, row: int, col: int, value) -> bool:
    """跳过合并单元格（只读），写入成功返回 True。"""
    from openpyxl.cell.cell import MergedCell

    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        return False
    cell.value = value
    return True


def _pick_enumerated_choice(option_text: str, situation: str) -> str | None:
    """对「1.xxx / 2.xxx」类选项，若选项文本本身可作为指标选择值则返回。"""
    opt = option_text.strip()
    if re.match(r"^\d+\.", opt):
        core = _option_core(opt)
        if core and core[:20] in situation:
            return opt.split(".", 1)[0].strip() + "." + opt.split(".", 1)[1][:80]
    return None


def fill_collection(
    template: Path,
    elements_path: Path,
    out_path: Path,
) -> dict:
    from openpyxl import load_workbook

    data = json.loads(elements_path.read_text(encoding="utf-8"))
    emap = _element_map(data)
    if not emap:
        raise ValueError("elements_extracted.json 中无要素")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, out_path)

    wb = load_workbook(out_path)
    ws = wb[wb.sheetnames[0]]

    from openpyxl.cell.cell import MergedCell

    for row in range(2, (ws.max_row or 0) + 1):
        c = ws.cell(row, 3).value
        if c is None or not str(c).strip():
            continue
        for col in (4, 5):
            cell = ws.cell(row, col)
            if not isinstance(cell, MergedCell):
                cell.value = None

    current_name: str | None = None
    filled_remarks = 0
    filled_choices = 0

    for row in range(2, (ws.max_row or 0) + 1):
        a = ws.cell(row, 1).value
        if a is not None and str(a).strip():
            current_name = str(a).strip()

        option = ws.cell(row, 3).value
        if not current_name or option is None or not str(option).strip():
            continue

        el = emap.get(current_name)
        if not el:
            for k, v in emap.items():
                if k in current_name or current_name in k:
                    el = v
                    break
        if not el:
            continue

        opt_str = str(option).strip()
        remark = _remark_for_element(el)
        cur_remark = ws.cell(row, 5).value
        if not cur_remark and _set_cell_value(ws, row, 5, remark):
            filled_remarks += 1

        choice_cell = ws.cell(row, 4)
        if isinstance(choice_cell, MergedCell):
            continue

        situation = el.get("situation") or ""
        logic = el.get("logic_trace") or ""

        if opt_str.startswith("□") or "□" in opt_str[:2]:
            yn = _infer_yes_no(opt_str, situation, logic)
            if yn and _set_cell_value(ws, row, 4, yn):
                filled_choices += 1
        elif re.match(r"^\d+\.", opt_str):
            picked = _pick_enumerated_choice(opt_str, situation)
            if picked and _set_cell_value(ws, row, 4, picked):
                filled_choices += 1
            elif (
                "信托" in current_name
                and "设立" in situation
                and "未设立" not in situation
                and opt_str.startswith("1.")
                and _set_cell_value(ws, row, 4, opt_str)
            ):
                filled_choices += 1

    wb.save(out_path)
    return {
        "filled_remarks": filled_remarks,
        "filled_choices": filled_choices,
        "elements": len(emap),
    }


def main() -> None:
    p = argparse.ArgumentParser(description="Case1 指标表兜底回填")
    p.add_argument("--template", required=True)
    p.add_argument("--elements", required=True)
    p.add_argument("--out", required=True)
    args = p.parse_args()

    cwd = Path.cwd().resolve()
    template = (cwd / args.template).resolve()
    elements = (cwd / args.elements).resolve()
    out = (cwd / args.out).resolve()
    for path in (template, elements):
        if not str(path).startswith(str(cwd)):
            raise SystemExit("路径必须位于当前工作目录之下")
    if not template.is_file():
        raise SystemExit(f"模板不存在：{template}")
    if not elements.is_file():
        raise SystemExit(f"要素 JSON 不存在：{elements}")

    catalog = cwd / "outputs" / "template_row_catalog.json"
    if catalog.is_file():
        raise SystemExit(
            "已存在 outputs/template_row_catalog.json，请续跑智能体按模板逐行填表，"
            "勿使用本应急脚本（见 SOP_信息收集表填报.md）。"
        )

    stats = fill_collection(template, elements, out)
    print(json.dumps({"ok": True, **stats}, ensure_ascii=False))


if __name__ == "__main__":
    main()
