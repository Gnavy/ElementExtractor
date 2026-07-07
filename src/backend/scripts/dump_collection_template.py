#!/usr/bin/env python3
"""
解析 Case1 信息收集表模板，输出逐行/逐组填表目录 template_row_catalog.json。

用法（cwd = extract_root）：
  python tools/dump_collection_template.py --src inputs/collection_template.xlsx
  python tools/dump_collection_template.py --src inputs/collection_template.xlsx --out outputs/template_row_catalog.json
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

HEADER_NAMES = frozenset({"指标名称", "指标解释", "指标选项", "指标选择", "备注"})
SECTION_MARKERS = frozenset({"客观部分指标", "主观部分指标"})
DATA_SOURCE_COL = 6  # 模板第 6 列：数据源优先顺序（F 列）

# 需联网查询当地政策/规划的指标关键词（catalog 标记 requires_tavily_search）
TAVILY_INDICATOR_KEYWORDS = (
    "个性化政策",
    "规划指标",
    "当地政策",
    "区域政策",
    "地方政策",
    "房地产政策",
    "城市政策",
    "限购",
    "限贷",
    "限售",
    "调控",
    "城市总体规划",
    "国土空间规划",
    "控规",
    "板块规划",
)

SOP_ELEMENT_HINTS = [
    "风险隔离措施（SPV层面）",
    "风险隔离措施",
    "本息支付情况",
    "增信措施约定",
    "SPV的分配顺序和清偿顺序",
    "项目违约后我司罚息利率超出固定收益率的倍数",
    "违约处置退出方式",
    "其他还款来源",
    "监管方案设置",
    "项目销售考核设置情况",
    "现金流入流出的确定性约定",
]


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _is_checkbox_option(text: str) -> bool:
    t = text.strip()
    return t.startswith("□") or t.startswith("☑") or t.startswith("✓")


def _is_numbered_option(text: str) -> bool:
    return bool(re.match(r"^\d+\.", text.strip()))


def _sop_hint(name: str) -> str | None:
    for hint in SOP_ELEMENT_HINTS:
        if hint in name or name in hint:
            return hint
    return None


def _needs_tavily_search(name: str, expl: str, data_source: str) -> bool:
    text = f"{name} {expl} {data_source}"
    return any(kw in text for kw in TAVILY_INDICATOR_KEYWORDS)


def _tavily_query_hint(name: str, expl: str) -> str:
    base = name.strip() or "当地政策"
    if expl.strip():
        return f"{base} {expl.strip()[:120]}"
    return base


def _parse_list_formula(formula: str | None) -> list[str] | None:
    if not formula:
        return None
    f = str(formula).strip()
    if f.startswith('"') and f.endswith('"') and "," in f:
        inner = f.strip('"')
        return [p.strip() for p in inner.split(",") if p.strip()]
    if f.startswith("$") and ":" in f:
        return None
    return None


def _collect_validations(ws) -> dict[str, list[str]]:
    """sqref -> allowed list (inline list only)."""
    out: dict[str, list[str]] = {}
    dvs = getattr(ws, "data_validations", None)
    if not dvs:
        return out
    for dv in dvs.dataValidation:
        choices = _parse_list_formula(getattr(dv, "formula1", None))
        if not choices:
            continue
        ref = str(dv.sqref)
        for part in ref.split():
            out[part] = choices
    return out


def _choices_for_row(row: int, validations: dict[str, list[str]]) -> list[str] | None:
    for ref, choices in validations.items():
        if ":" in ref:
            a, b = ref.split(":", 1)
            m = re.match(r"([A-Z]+)(\d+)", a)
            n = re.match(r"([A-Z]+)(\d+)", b)
            if not m or not n:
                continue
            if m.group(1) != "D" or n.group(1) != "D":
                continue
            r1, r2 = int(m.group(2)), int(n.group(2))
            if r1 <= row <= r2:
                return choices
        else:
            m = re.match(r"D(\d+)", ref)
            if m and int(m.group(1)) == row:
                return choices
    return None


def _infer_choice_mode(rows: list[dict], group_validations: list[str] | None) -> str:
    has_cb = any(_is_checkbox_option(r["option_text"]) for r in rows)
    has_num = any(_is_numbered_option(r["option_text"]) for r in rows)
    if has_cb and not has_num:
        return "yes_no"
    if has_num:
        return "exclusive"
    if group_validations == ["是", "否"]:
        return "yes_no"
    if group_validations and len(group_validations) > 2:
        return "exclusive"
    return "yes_no" if has_cb else "exclusive" if has_num else "unknown"


def parse_sheet(ws, sheet_name: str) -> list[dict]:
    validations = _collect_validations(ws)
    max_row = ws.max_row or 0
    current_section = "主观部分指标"
    current_name = ""
    current_expl = ""
    current_data_source_priority = ""
    section_groups: dict[str, list[dict]] = {}
    pending_rows: list[dict] = []

    def flush_group() -> None:
        nonlocal pending_rows, current_name, current_expl, current_section, current_data_source_priority
        if not pending_rows or not current_name:
            pending_rows = []
            return
        row_nums = [r["row"] for r in pending_rows]
        group_vals: list[str] | None = None
        for r in row_nums:
            c = _choices_for_row(r, validations)
            if c:
                group_vals = c
                break
        mode = _infer_choice_mode(pending_rows, group_vals)
        allowed = group_vals
        if mode == "yes_no":
            allowed = ["是", "否"]
        elif mode == "exclusive" and not allowed:
            allowed = [r["option_text"] for r in pending_rows if r["option_text"]]

        for r in pending_rows:
            row_allowed = _choices_for_row(r["row"], validations) or (
                allowed if mode == "yes_no" else None
            )
            r["allowed_choices"] = row_allowed or (
                allowed if mode == "exclusive" else ["是", "否"]
            )

        group_data_source = ""
        for r in pending_rows:
            ds = (r.get("data_source_priority") or "").strip()
            if ds:
                group_data_source = ds
                break

        grp = {
            "indicator_name": current_name,
            "indicator_explanation": current_expl,
            "choice_mode": mode,
            "sop_element_hint": _sop_hint(current_name),
            "allowed_choices": allowed,
            "data_source_priority": group_data_source or current_data_source_priority,
            "requires_tavily_search": _needs_tavily_search(
                current_name,
                current_expl,
                group_data_source or current_data_source_priority,
            ),
            "tavily_query_hint": _tavily_query_hint(current_name, current_expl),
            "rows": pending_rows,
        }
        section_groups.setdefault(current_section, []).append(grp)
        pending_rows = []

    for row in range(1, max_row + 1):
        a = _cell_str(ws.cell(row, 1).value)
        b = _cell_str(ws.cell(row, 2).value)
        c = _cell_str(ws.cell(row, 3).value)
        f = _cell_str(ws.cell(row, DATA_SOURCE_COL).value)

        if a in HEADER_NAMES and c in HEADER_NAMES:
            continue
        if a in SECTION_MARKERS:
            flush_group()
            current_section = a
            current_name = ""
            current_expl = ""
            current_data_source_priority = ""
            continue
        if a == "指标名称" or (a == "" and b == "" and c == ""):
            continue

        if a:
            flush_group()
            current_name = a
            current_expl = b if b else current_expl
            if f:
                current_data_source_priority = f
        elif b and not current_expl:
            current_expl = b
            if f and not current_data_source_priority:
                current_data_source_priority = f

        if not c:
            continue
        if not current_name:
            continue

        pending_rows.append(
            {
                "row": row,
                "option_text": c,
                "choice_column": "D",
                "remark_column": "E",
                "data_source_priority": f or current_data_source_priority,
            }
        )

    flush_group()

    if not section_groups:
        return []

    order = ["主观部分指标", "客观部分指标"]
    keys = [k for k in order if k in section_groups]
    keys += [k for k in section_groups if k not in keys]
    return [
        {"section_title": sec, "indicator_groups": section_groups[sec]}
        for sec in keys
    ]


def parse_workbook(src: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(src, read_only=False, data_only=True)
    sheets_out = []
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            sections = parse_sheet(ws, name)
            if sections:
                sheets_out.append({"name": name, "sections": sections})
    finally:
        wb.close()
    return {
        "source": str(src.name),
        "fill_rules": {
            "yes_no": "指标选项以□开头：每行独立填「是」或「否」",
            "exclusive": "组内多行1./2./3.选项：仅一行D列填入选项全文，其余留空",
            "remark": "备注一句话说明理由",
        },
        "sheets": sheets_out,
    }


def main() -> None:
    p = argparse.ArgumentParser(description="导出信息收集表行目录 JSON")
    p.add_argument("--src", required=True, help="如 inputs/collection_template.xlsx")
    p.add_argument(
        "--out",
        default="outputs/template_row_catalog.json",
        help="输出路径（相对 cwd）",
    )
    args = p.parse_args()

    cwd = Path.cwd().resolve()
    src = (cwd / args.src).resolve()
    out = (cwd / args.out).resolve()
    if not str(src).startswith(str(cwd)):
        raise SystemExit("路径必须位于当前工作目录之下")
    if not src.is_file():
        raise SystemExit(f"文件不存在：{src}")

    data = parse_workbook(src)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"ok": True, "out": str(args.out), "sheets": len(data["sheets"])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
