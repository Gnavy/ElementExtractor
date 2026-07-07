#!/usr/bin/env python3
"""
Case2 模板抽象为 fill_plan（整表填报计划）。

每个 item = 一个待填单元（财务报表通常为「一行科目 + 多列金额」；
指标表通常为「一行选项 + 选择格 + 备注格」）。

输出 outputs/case2_fill_schema.json（schema_version=2）
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

SCHEMA_VERSION = 2


def _s(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _parse_inline_list(formula: str | None) -> list[str] | None:
    if not formula:
        return None
    f = str(formula).strip()
    if f.startswith('"') and f.endswith('"') and "," in f:
        return [x.strip() for x in f.strip('"').split(",") if x.strip()]
    return None


def _collect_dv(ws) -> dict[int, list[str]]:
    out: dict[int, list[str]] = {}
    dvs = getattr(ws, "data_validations", None)
    if not dvs:
        return out
    for dv in dvs.dataValidation:
        choices = _parse_inline_list(getattr(dv, "formula1", None))
        if not choices:
            continue
        for part in str(dv.sqref).split():
            if ":" in part:
                a, b = part.split(":", 1)
                ma = re.match(r"([A-Z]+)(\d+)", a)
                mb = re.match(r"([A-Z]+)(\d+)", b)
                if not ma or not mb:
                    continue
                if ma.group(1) != "D" or mb.group(1) != "D":
                    continue
                r1, r2 = int(ma.group(2)), int(mb.group(2))
                for r in range(r1, r2 + 1):
                    out[r] = choices
            else:
                m = re.match(r"D(\d+)", part)
                if m:
                    out[int(m.group(1))] = choices
    return out


def _infer_value_type(text: str, allowed_values: list[str] | None) -> str:
    if allowed_values:
        if set(allowed_values) == {"是", "否"} or set(allowed_values) == {"否", "是"}:
            return "boolean"
        return "enum"
    t = text.strip()
    if re.search(r"\d{4}[-/年]", t):
        return "date"
    if re.search(r"\d", t):
        return "number_or_text"
    return "text"


def _empty_field(
    cell: str,
    *,
    column_label: str = "",
    value_type: str = "text",
    allowed_values: list[str] | None = None,
) -> dict:
    return {
        "cell": cell,
        "column_label": column_label,
        "value_type": value_type,
        "allowed_values": allowed_values or [],
        "value": None,
    }


def _sheet_looks_financial(ws) -> bool:
    for r in range(1, 10):
        b = _s(ws.cell(r, 2).value).replace(" ", "")
        if "科目" in b:
            return True
    return False


def _read_column_headers(ws, row: int = 3) -> dict[str, str]:
    headers: dict[str, str] = {}
    for col_idx, letter in ((3, "C"), (4, "D"), (5, "E"), (6, "F")):
        h = _s(ws.cell(row, col_idx).value)
        if h:
            headers[letter] = h
    return headers


def _build_financial_items(ws, sheet_name: str) -> tuple[list[dict], dict[str, str]]:
    from openpyxl.cell.cell import MergedCell

    column_headers = _read_column_headers(ws)
    items: list[dict] = []
    max_row = min(ws.max_row or 0, 124)

    for r in range(4, max_row + 1):
        label = _s(ws.cell(r, 2).value)
        if not label:
            continue
        fields: dict[str, dict] = {}
        for col_idx, letter in ((3, "C"), (4, "D"), (5, "E"), (6, "F")):
            cell = ws.cell(r, col_idx)
            if isinstance(cell, MergedCell):
                continue
            if cell.data_type == "f" or (
                isinstance(cell.value, str) and str(cell.value).startswith("=")
            ):
                continue
            vt = "date" if r == 4 or "日期" in label else "number"
            fields[letter] = _empty_field(
                cell.coordinate,
                column_label=column_headers.get(letter, ""),
                value_type=vt,
            )
        if not fields:
            continue
        meaning = label
        if column_headers:
            cols = "/".join(column_headers.values())
            meaning = f"{label}（{cols}）"
        items.append(
            {
                "item_id": f"{sheet_name}:r{r}",
                "row": r,
                "label": label,
                "meaning": meaning,
                "data_source_priority": "",
                "fields": fields,
                "confidence": None,
                "reason_one_line": None,
                "evidence_refs": [],
            }
        )
    return items, column_headers


def _indicator_dim_to_item(
    sheet_name: str,
    indicator_name: str,
    indicator_expl: str,
    data_source_priority: str,
    dim: dict,
) -> dict:
    fields: dict[str, dict] = {}
    target_cells = dim.get("target_cells") or []
    remark_cells = dim.get("remark_cells") or []
    if target_cells:
        fields["choice"] = _empty_field(
            target_cells[0],
            value_type=dim.get("value_type") or "text",
            allowed_values=dim.get("allowed_values"),
        )
    if remark_cells:
        fields["remark"] = _empty_field(remark_cells[0], value_type="text")
    return {
        "item_id": f"{sheet_name}:r{dim['row']}",
        "row": dim["row"],
        "label": indicator_name,
        "meaning": dim.get("prompt") or indicator_name,
        "option_text": dim.get("prompt") or "",
        "indicator_explanation": indicator_expl,
        "data_source_priority": dim.get("data_source_priority") or data_source_priority,
        "fields": fields,
        "confidence": None,
        "reason_one_line": None,
        "evidence_refs": [],
    }


def build_schema(src: Path) -> tuple[dict, dict]:
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell

    wb = load_workbook(src, read_only=False, data_only=False)
    try:
        sheets_out: list[dict] = []
        catalog_sheets: list[dict] = []
        template_instructions: list[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            max_row = ws.max_row or 0
            dv_map = _collect_dv(ws)

            if sheet_name in ("填表说明",) or _sheet_looks_financial(ws) is False:
                for r in range(1, min(12, max_row + 1)):
                    t = _s(ws.cell(r, 2).value)
                    if t and len(t) > 8:
                        template_instructions.append(t)

            items: list[dict] = []
            sections: list[dict] = [{"section_title": "default", "indicator_groups": []}]
            sec_groups = sections[0]["indicator_groups"]

            if _sheet_looks_financial(ws):
                items, column_headers = _build_financial_items(ws, sheet_name)
                sheets_out.append(
                    {
                        "sheet": sheet_name,
                        "column_headers": column_headers,
                        "items": items,
                    }
                )
                catalog_sheets.append({"name": sheet_name, "sections": sections})
                continue

            cur_name = ""
            cur_expl = ""
            cur_ds = ""
            cur_dims: list[dict] = []

            def flush_indicator() -> None:
                nonlocal cur_name, cur_expl, cur_ds, cur_dims
                if not cur_name or not cur_dims:
                    cur_dims = []
                    return
                for d in cur_dims:
                    items.append(
                        _indicator_dim_to_item(
                            sheet_name, cur_name, cur_expl, cur_ds, d
                        )
                    )
                has_numbered = any(
                    re.match(r"^\d+\.", d["prompt"]) for d in cur_dims
                )
                has_checkbox = any(
                    d["prompt"].lstrip().startswith(("□", "☑", "✓"))
                    for d in cur_dims
                )
                mode = (
                    "exclusive"
                    if has_numbered
                    else "yes_no"
                    if has_checkbox
                    else "unknown"
                )
                sec_groups.append(
                    {
                        "indicator_name": cur_name,
                        "indicator_explanation": cur_expl,
                        "choice_mode": mode,
                        "allowed_choices": cur_dims[0].get("allowed_values")
                        if cur_dims
                        else None,
                        "data_source_priority": cur_ds,
                        "rows": [
                            {
                                "row": d["row"],
                                "option_text": d["prompt"],
                                "choice_column": "D",
                                "remark_column": "E",
                                "allowed_choices": d.get("allowed_values") or [],
                                "data_source_priority": d.get("data_source_priority")
                                or cur_ds,
                            }
                            for d in cur_dims
                        ],
                    }
                )
                cur_dims = []

            for r in range(1, max_row + 1):
                a = _s(ws.cell(r, 1).value)
                b = _s(ws.cell(r, 2).value)
                c = _s(ws.cell(r, 3).value)
                f = _s(ws.cell(r, 6).value)

                if a == "指标名称" and c == "指标选项":
                    continue
                if a in ("主观部分指标", "客观部分指标"):
                    flush_indicator()
                    continue
                if a:
                    flush_indicator()
                    cur_name = a
                    cur_expl = b or cur_expl
                    cur_ds = f or cur_ds
                elif b and not cur_expl:
                    cur_expl = b
                    if f and not cur_ds:
                        cur_ds = f

                if not c or not cur_name:
                    continue

                d_cell = ws.cell(r, 4)
                e_cell = ws.cell(r, 5)
                if isinstance(d_cell, MergedCell):
                    continue
                allowed = dv_map.get(r)
                cur_dims.append(
                    {
                        "row": r,
                        "prompt": c,
                        "value_type": _infer_value_type(c, allowed),
                        "allowed_values": allowed or [],
                        "target_cells": [d_cell.coordinate],
                        "remark_cells": [e_cell.coordinate]
                        if not isinstance(e_cell, MergedCell)
                        else [],
                        "data_source_priority": f or cur_ds,
                    }
                )

            flush_indicator()
            sheets_out.append({"sheet": sheet_name, "items": items})
            catalog_sheets.append({"name": sheet_name, "sections": sections})

    finally:
        wb.close()

    item_count = sum(len(s.get("items") or []) for s in sheets_out)
    field_count = sum(
        len(item.get("fields") or {})
        for s in sheets_out
        for item in s.get("items") or []
    )
    mode = "fill_plan"
    if any((s.get("items") or []) and s.get("column_headers") for s in sheets_out):
        mode = "fill_plan_financial"

    schema = {
        "schema_version": SCHEMA_VERSION,
        "source": src.name,
        "mode": mode,
        "template_instructions": template_instructions[:20],
        "item_count": item_count,
        "field_count": field_count,
        "sheets": sheets_out,
    }
    catalog = {"source": src.name, "sheets": catalog_sheets}
    return schema, catalog


def main() -> None:
    p = argparse.ArgumentParser(description="导出 Case2 fill_plan")
    p.add_argument("--src", required=True)
    p.add_argument("--out", default="outputs/case2_fill_schema.json")
    p.add_argument("--catalog-out", default="outputs/template_row_catalog.json")
    args = p.parse_args()

    cwd = Path.cwd().resolve()
    src = (cwd / args.src).resolve()
    out = (cwd / args.out).resolve()
    cat = (cwd / args.catalog_out).resolve()
    if not str(src).startswith(str(cwd)) or not src.is_file():
        raise SystemExit("模板路径无效")
    if not str(out).startswith(str(cwd)) or not str(cat).startswith(str(cwd)):
        raise SystemExit("输出路径必须在当前目录内")

    schema, catalog = build_schema(src)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(schema, ensure_ascii=False, indent=2), encoding="utf-8")
    cat.parent.mkdir(parents=True, exist_ok=True)
    cat.write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "ok": True,
                "schema_out": args.out,
                "catalog_out": args.catalog_out,
                "item_count": schema.get("item_count"),
                "field_count": schema.get("field_count"),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
