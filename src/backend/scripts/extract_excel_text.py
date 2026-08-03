#!/usr/bin/env python3
"""
提取 Excel 内容为文本：JSON 索引 + ocr_text/<相对路径>.md（逐 sheet 的 markdown 表格）。

跳过 inputs/——待填模板由 dump_collection_template*.py 负责，那是结构不是内容。

用法（cwd = extract_root）：
  python tools/extract_excel_text.py --src . --out outputs/excel_text_index.json
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

# 单 sheet 导出上限，避免超大表撑爆语料
MAX_ROWS = 500
MAX_COLS = 40
SUPPORTED = {".xlsx", ".xlsm"}
LEGACY = {".xls"}
SKIP_DIRS = {"outputs", "inputs", "tools", "ocr_text"}


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text.replace("|", "\\|").replace("\n", " ")


def _sheet_to_markdown(ws) -> tuple[str, int, int]:
    """返回（markdown 表格, 实际导出行数, 实际导出列数）。"""
    rows = []
    max_col = 0
    for row in ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True):
        cells = [_cell_text(v) for v in row]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            continue
        max_col = max(max_col, len(cells))
        rows.append(cells)
    if not rows:
        return "", 0, 0
    lines = []
    for index, cells in enumerate(rows):
        padded = cells + [""] * (max_col - len(cells))
        lines.append("| " + " | ".join(padded) + " |")
        if index == 0:
            lines.append("|" + "---|" * max_col)
    return "\n".join(lines), len(rows), max_col


def _legacy_sheet_to_markdown(sheet) -> tuple[str, int, int]:
    """xlrd 的 sheet → markdown，口径与 openpyxl 那条一致。"""
    rows = []
    max_col = 0
    for r in range(min(sheet.nrows, MAX_ROWS)):
        cells = [_cell_text(v) for v in sheet.row_values(r)[:MAX_COLS]]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            continue
        max_col = max(max_col, len(cells))
        rows.append(cells)
    if not rows:
        return "", 0, 0
    lines = []
    for index, cells in enumerate(rows):
        padded = cells + [""] * (max_col - len(cells))
        lines.append("| " + " | ".join(padded) + " |")
        if index == 0:
            lines.append("|" + "---|" * max_col)
    return "\n".join(lines), len(rows), max_col


def extract_legacy_workbook(path: Path) -> dict:
    """老的 .xls 二进制格式，openpyxl 读不了，走 xlrd。"""
    import xlrd  # 延迟导入：未装时只影响 .xls，不影响 xlsx

    book = xlrd.open_workbook(str(path))
    sheets = []
    blocks = []
    for sheet in book.sheets():
        table, n_rows, n_cols = _legacy_sheet_to_markdown(sheet)
        sheets.append({"name": sheet.name, "rows": n_rows, "cols": n_cols})
        if table:
            blocks.append(f"## {sheet.name}\n\n{table}")
    return {"file": path.as_posix(), "sheets": sheets, "text": "\n\n".join(blocks)}


def extract_workbook(path: Path) -> dict:
    # data_only=True 取公式缓存值
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        sheets = []
        blocks = []
        for ws in wb.worksheets:
            table, n_rows, n_cols = _sheet_to_markdown(ws)
            sheets.append({"name": ws.title, "rows": n_rows, "cols": n_cols})
            if table:
                blocks.append(f"## {ws.title}\n\n{table}")
        return {
            "file": path.as_posix(),
            "sheets": sheets,
            "text": "\n\n".join(blocks),
        }
    finally:
        wb.close()


def main() -> None:
    p = argparse.ArgumentParser(description="提取 Excel 内容为文本")
    p.add_argument("--src", required=True, help="Excel 文件或目录")
    p.add_argument("--out", required=True, help="输出 json")
    args = p.parse_args()

    cwd = Path.cwd().resolve()
    src = (cwd / args.src).resolve()
    out = (cwd / args.out).resolve()
    if not str(src).startswith(str(cwd)) or not str(out).startswith(str(cwd)):
        raise SystemExit("路径必须位于当前工作目录下")

    files = [src] if src.is_file() else sorted(src.rglob("*"))
    docs: list[dict] = []
    skipped: list[dict] = []
    ocr_root = cwd / "ocr_text"

    for f in files:
        if not f.is_file():
            continue
        suffix = f.suffix.lower()
        if suffix not in SUPPORTED and suffix not in LEGACY:
            continue
        rel_parts = f.relative_to(cwd).parts
        if rel_parts and rel_parts[0] in SKIP_DIRS:
            continue
        if f.name.startswith("~$"):  # Office 临时文件
            continue
        rel = f.relative_to(cwd)

        try:
            if suffix in LEGACY:
                info = extract_legacy_workbook(f)
            else:
                info = extract_workbook(f)
        except ImportError:
            skipped.append({"file": rel.as_posix(), "reason": ".xls 需要 xlrd，当前环境未安装"})
            continue
        except Exception as exc:  # noqa: BLE001 — 单个文件坏掉不拖垮整批
            skipped.append({"file": rel.as_posix(), "reason": f"读取失败: {exc}"[:200]})
            continue

        body = info.pop("text", "") or ""
        md_path = ocr_root / f"{rel.as_posix()}.md"
        md_path.parent.mkdir(parents=True, exist_ok=True)
        md_path.write_text(
            f"# {rel.as_posix()}\n\n" + body + ("\n" if body else ""), encoding="utf-8"
        )
        info["file"] = rel.as_posix()
        info["md_path"] = md_path.relative_to(cwd).as_posix()
        info["char_count"] = len(body)
        info["preview"] = body[:800]
        docs.append(info)

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(
        json.dumps({"workbooks": docs, "skipped": skipped}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(out.relative_to(cwd).as_posix())


if __name__ == "__main__":
    main()
