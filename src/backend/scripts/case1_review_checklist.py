#!/usr/bin/env python3
"""生成 Case1 人工复核清单。

把散在 validation_report.json 里的 error/warning 按行汇总，配上该行实际填了什么、
备注写了什么，供人工逐条核对。**只读产物，不改任何东西，也不参与任务成败判定。**

与运行链路解耦：不挂进 LangGraph，任务跑完后单独执行。

用法：
  python scripts/case1_review_checklist.py --extract-root <任务目录> [--out <文件>]
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

# 问题类型 → 复核时该怎么看
_HOW_TO_CHECK = {
    "未填指标选择": "模型未给出结论。核对材料后人工补填，或确认该行确实无从判断",
    "结论与选择相反": "备注的理由与所填选项矛盾，二者必有一错。以材料为准重新判定",
    "备注混入推演过程": "备注里留了模型的思考草稿，需改写为一句判断要点 + 引用。结论本身未必错",
    "缺少源文件名与原文引用": "无法溯源，须回材料确认结论是否有据，并补齐引用",
    "带豁免条款": "模板对该条目写了豁免情形，需人工确认本项目是否属于该情形",
    "未选最优项": "同上，该组带豁免条款，确认是否应按满分口径",
    "指标选择须为是/否": "取值非法，必须修正",
    "未选择任何选项": "互斥组必须恰选一项，须补填",
    "有多行填了指标选择": "互斥组只能选一项，须删除多余项",
}


def _how_to_check(msg: str) -> str:
    for key, tip in _HOW_TO_CHECK.items():
        if key in msg:
            return tip
    return "核对材料确认"


def _row_of(msg: str) -> int | None:
    m = re.search(r"行(\d+)", msg)
    return int(m.group(1)) if m else None


def build_checklist(extract_root: Path) -> str:
    outputs = extract_root / "outputs"
    report_path = outputs / "validation_report.json"
    filled_path = outputs / "collection_filled.xlsx"
    if not report_path.is_file():
        return "# Case1 人工复核清单\n\n未找到 validation_report.json，无法生成。\n"

    report = json.loads(report_path.read_text(encoding="utf-8"))
    errors = report.get("errors") or []
    warnings = report.get("warnings") or []

    # 取实际填报值，便于人工对照，取不到不影响清单生成
    cells: dict[int, tuple[str, str]] = {}
    if filled_path.is_file():
        try:
            from openpyxl import load_workbook

            wb = load_workbook(filled_path, read_only=True, data_only=True)
            try:
                ws = wb[wb.sheetnames[0]]
                for r in range(1, ws.max_row + 1):
                    d = ws.cell(r, 4).value
                    e = ws.cell(r, 5).value
                    if d not in (None, "") or e not in (None, ""):
                        cells[r] = (str(d or "").strip(), str(e or "").strip())
            finally:
                wb.close()
        except Exception:  # noqa: BLE001 - 清单是辅助产物，读不到就降级
            cells = {}

    lines = ["# Case1 人工复核清单", ""]
    stats = report.get("stats") or {}
    lines.append(
        f"校验结论：{'通过' if report.get('passed') else '未通过'}　|　"
        f"错误 {len(errors)}　警告 {len(warnings)}　|　"
        f"是否行 {stats.get('yes_no_rows', '?')}　互斥组 {stats.get('exclusive_groups', '?')}"
    )
    lines.append("")
    lines.append(
        "> 校验通过只代表格式与结构合规，**不代表填报内容正确**。"
        "下列条目是机器能发现的问题，机器发现不了的错误仍需抽查。"
    )
    lines.append("")

    for title, items, note in (
        ("## 必须处理（error）", errors, "不处理则任务判定为失败"),
        ("## 建议复核（warning）", warnings, "不阻塞任务，但很可能是错的"),
    ):
        lines.append(title)
        lines.append("")
        if not items:
            lines.append("无。")
            lines.append("")
            continue
        lines.append(f"_{note}_")
        lines.append("")
        for msg in items:
            row = _row_of(msg)
            lines.append(f"- **{msg}**")
            lines.append(f"  - 怎么看：{_how_to_check(msg)}")
            if row and row in cells:
                choice, remark = cells[row]
                lines.append(f"  - 当前填报：`{choice[:60] or '（空）'}`")
                if remark:
                    lines.append(f"  - 当前备注：{remark[:120]}")
            lines.append("")

    lines.append("## 机器查不出来的，请抽查")
    lines.append("")
    lines.append(
        "- **引用真实但结论相反**：引用原文摘录是真的，判断要点却与它方向相反。"
        "校验层只能检出显式的自我否定说法，这类需要人读原文比对"
    )
    lines.append(
        "- **数值比较与计数**：模型可能列出正确数据却比较错、数错。"
        "凡涉及阈值、个数、比例的行，建议按引用原文复算一遍"
    )
    lines.append(
        "- **同份材料内证据覆盖不全**：结论只采信了材料的一部分，"
        "另一部分与之矛盾的事实未被引用"
    )
    lines.append("")
    return "\n".join(lines)


def main() -> None:
    p = argparse.ArgumentParser(description="生成 Case1 人工复核清单")
    p.add_argument("--extract-root", required=True)
    p.add_argument("--out", default="", help="默认写到 <extract-root>/outputs/case1_review_checklist.md")
    args = p.parse_args()

    root = Path(args.extract_root)
    text = build_checklist(root)
    out = Path(args.out) if args.out else root / "outputs" / "case1_review_checklist.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")
    print(f"已生成：{out}")
    print(text[:600])


if __name__ == "__main__":
    main()
