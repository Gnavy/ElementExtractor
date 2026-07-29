#!/usr/bin/env python3
"""
将 Case2 filled fill_plan / legacy schema 回填到 xlsx。
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Iterator

_SCRIPT_DIR = Path(__file__).resolve().parent
if str(_SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPT_DIR))

from ocr_confidence_export import default_confidence_threshold  # noqa: E402
from ocr_confidence_match import (  # noqa: E402
    _ocr_md_refs,
    is_low_ocr_confidence,
    load_sidecar_index,
)


def _norm(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _target_coords(dim: dict) -> list[str]:
    cells = dim.get("target_cells") or []
    if not cells and dim.get("target_cell"):
        cells = [dim["target_cell"]]
    return [str(c).strip() for c in cells if c]


def _remark_coords(dim: dict) -> list[str]:
    cells = dim.get("remark_cells") or []
    if not cells and dim.get("remark_cell"):
        cells = [dim["remark_cell"]]
    return [str(c).strip() for c in cells if c]


def _iter_fill_plan_writes(
    data: dict,
) -> Iterator[tuple[str, str, Any, list[Any], str]]:
    """yield (sheet, cell_coord, value, evidence_refs, value_type)"""
    for sh in data.get("sheets") or []:
        sheet = sh.get("sheet") or sh.get("name")
        if not sheet:
            continue
        for item in sh.get("items") or []:
            reason = _norm(item.get("reason_one_line"))
            evidence_refs = item.get("evidence_refs") or []
            for _key, field in (item.get("fields") or {}).items():
                if not isinstance(field, dict):
                    continue
                cell = field.get("cell")
                if not cell:
                    continue
                if "value" not in field:
                    continue
                yield (
                    sheet,
                    cell,
                    field.get("value"),
                    evidence_refs,
                    str(field.get("value_type") or ""),
                )
            remark = (item.get("fields") or {}).get("remark")
            if isinstance(remark, dict) and remark.get("cell") and reason:
                yield sheet, remark["cell"], reason, evidence_refs, "text"


def _has_ocr_evidence(data: dict) -> bool:
    for sh in data.get("sheets") or []:
        for item in sh.get("items") or []:
            if _ocr_md_refs(item.get("evidence_refs") or []):
                return True
    return False


def _iter_legacy_dimensions(data: dict) -> Iterator[tuple[str, dict, dict]]:
    for ss in data.get("subject_sets") or []:
        sheet = ss.get("sheet")
        for subj in ss.get("subjects") or []:
            for dim in subj.get("dimensions") or []:
                yield sheet, subj, dim
    for sh in data.get("sheets") or []:
        if sh.get("items"):
            continue
        sheet = sh.get("sheet") or sh.get("name")
        for subj in sh.get("subjects") or []:
            for dim in subj.get("dimensions") or []:
                yield sheet, subj, dim


def _is_empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip().lower() in {"", "null", "none", "nil", "n/a", "na"}
    return False


def _coerce_value(value: Any, value_type: str) -> Any:
    if _is_empty(value):
        return None
    if not isinstance(value, str):
        return value
    text = value.strip()
    if value_type == "number":
        negative = text.startswith("(") and text.endswith(")")
        number = text.strip("()").replace(",", "").replace("，", "").strip()
        if re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", number):
            parsed = float(number)
            if negative:
                parsed = -parsed
            return int(parsed) if parsed.is_integer() else parsed
    if value_type == "date":
        normalized = (
            text.replace("年", "-").replace("月", "-").replace("日", "")
            .replace("/", "-")
            .replace(".", "-")
        )
        for fmt in ("%Y-%m-%d", "%Y-%m", "%Y"):
            try:
                parsed = datetime.strptime(normalized, fmt)
                if fmt == "%Y":
                    parsed = parsed.replace(month=12, day=31)
                elif fmt == "%Y-%m":
                    parsed = parsed.replace(day=1)
                return parsed
            except ValueError:
                continue
    return text


def main() -> None:
    p = argparse.ArgumentParser(description="Case2 fill_plan 回填 xlsx")
    p.add_argument("--template", required=True)
    p.add_argument("--filled-schema", required=True)
    p.add_argument("--out", default="outputs/collection_filled.xlsx")
    p.add_argument("--report", default="outputs/backfill_report.json")
    args = p.parse_args()

    cwd = Path.cwd().resolve()
    template = (cwd / args.template).resolve()
    filled = (cwd / args.filled_schema).resolve()
    out = (cwd / args.out).resolve()
    report = (cwd / args.report).resolve()
    for pth in (template, filled, out, report):
        if not str(pth).startswith(str(cwd)):
            raise SystemExit("路径必须位于当前工作目录之下")
    if not template.is_file() or not filled.is_file():
        raise SystemExit("模板或 filled schema 不存在")

    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill

    data = json.loads(filled.read_text(encoding="utf-8"))
    ocr_threshold = default_confidence_threshold()
    sidecar_index = load_sidecar_index(cwd)
    sidecar_count = len(sidecar_index)
    ocr_marked: list[dict[str, Any]] = []
    ocr_low_confidence_cells = 0
    yellow_fill = PatternFill(fill_type="solid", fgColor="FFFF00")

    if sidecar_count == 0 and _has_ocr_evidence(data):
        warnings_pre: list[str] = [
            "未找到 OCR sidecar（*.ocr_cells.json）；低置信度标黄已跳过。"
            "请重新执行 OCR 以生成 sidecar。"
        ]
    else:
        warnings_pre = []

    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, out)
    wb = load_workbook(out, read_only=False, data_only=False)
    written = 0
    target_cells = 0
    empty_cells = 0
    skipped = 0
    warnings: list[str] = list(warnings_pre)
    use_plan = any((sh.get("items") or []) for sh in data.get("sheets") or [])

    def write_cell(
        ws,
        coord: str,
        value: Any,
        evidence_refs: list[Any] | None = None,
        value_type: str = "",
    ) -> None:
        nonlocal written, target_cells, empty_cells, skipped, ocr_low_confidence_cells
        target_cells += 1
        value = _coerce_value(value, value_type)
        if value is None:
            empty_cells += 1
            return
        c = ws[coord]
        if isinstance(c, MergedCell):
            skipped += 1
            return
        if c.data_type == "f" or (
            isinstance(c.value, str) and str(c.value).startswith("=")
        ):
            skipped += 1
            warnings.append(f"{ws.title}:{coord} 为公式，已跳过")
            return
        c.value = value
        if value_type == "date":
            c.number_format = "yyyy-mm"
        written += 1

        if not sidecar_index or evidence_refs is None:
            return
        if value is None or value == "":
            return

        low, min_conf, md_ref = is_low_ocr_confidence(
            value, evidence_refs, sidecar_index, threshold=ocr_threshold
        )
        if not low or min_conf is None:
            return

        c.fill = yellow_fill
        pct = int(round(min_conf * 100))
        source = md_ref or "ocr_text"
        c.comment = Comment(
            f"OCR置信度 {pct}%，来源：{source}",
            "ElementExtractor",
        )
        ocr_low_confidence_cells += 1
        ocr_marked.append(
            {
                "sheet": ws.title,
                "cell": coord,
                "value": value,
                "ocr_confidence": round(min_conf, 4),
                "evidence_md": md_ref,
            }
        )

    try:
        if use_plan:
            for sheet, coord, value, evidence_refs, value_type in _iter_fill_plan_writes(
                data
            ):
                if sheet not in wb.sheetnames:
                    warnings.append(f"sheet 不存在: {sheet}")
                    skipped += 1
                    continue
                write_cell(wb[sheet], coord, value, evidence_refs, value_type)
        else:
            for sheet, subj, dim in _iter_legacy_dimensions(data):
                if not sheet or sheet not in wb.sheetnames:
                    warnings.append(f"sheet 不存在: {sheet}")
                    skipped += 1
                    continue
                ws = wb[sheet]
                value = dim.get("value")
                reason = _norm(dim.get("reason_one_line"))
                allowed = dim.get("allowed_values") or []
                sval = _norm(value)
                if allowed and sval and sval not in allowed:
                    warnings.append(
                        f"{sheet}:{subj.get('subject_name')}:{dim.get('dimension_key')} "
                        f"值 {sval!r} 不在候选中"
                    )
                for coord in _target_coords(dim):
                    if value is None and not reason:
                        skipped += 1
                        continue
                    write_cell(ws, coord, value)
                for coord in _remark_coords(dim):
                    if reason:
                        write_cell(ws, coord, reason)
    finally:
        wb.save(out)
        wb.close()

    rep = {
        "ok": True,
        "format": "fill_plan" if use_plan else "legacy",
        "target_cells": target_cells,
        "written_cells": written,
        "empty_cells": empty_cells,
        "skipped_cells": skipped,
        "ocr_sidecar_count": sidecar_count,
        "ocr_confidence_threshold": ocr_threshold,
        "ocr_low_confidence_cells": ocr_low_confidence_cells,
        "ocr_marked": ocr_marked[:100],
        "warnings": warnings[:50],
    }
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(json.dumps(rep, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(rep, ensure_ascii=False))


if __name__ == "__main__":
    main()
