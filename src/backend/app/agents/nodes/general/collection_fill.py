from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from app.agents.llm import structured_llm
from app.agents.prompts import general as prompts
from app.agents.schemas.extraction import CollectionFillPlan
from app.agents.tools.context import collect_ocr_snippets
from app.agents.tools.script_runner import run_tool_script


def _col_to_index(col: Any) -> int | None:
    if isinstance(col, int):
        return col
    if isinstance(col, str):
        s = col.strip().upper()
        if s.isdigit():
            return int(s)
        # Excel letter
        n = 0
        for ch in s:
            if "A" <= ch <= "Z":
                n = n * 26 + (ord(ch) - 64)
            else:
                return None
        return n or None
    return None


def maybe_mark_collection_node(state: dict[str, Any]) -> dict[str, Any]:
    meta = state.get("meta") or {}
    path = meta.get("collection_template_path")
    need = bool(path)
    return {
        "need_collection_fill": need,
        "collection_template_path": path,
        "log_lines": [
            "检测到信息收集表，将执行填报" if need else "无需信息收集表填报"
        ],
    }


def collection_fill_node(state: dict[str, Any]) -> dict[str, Any]:
    root = Path(state["extract_root"])
    meta = state.get("meta") or {}
    task_id = state.get("task_id") or meta.get("task_id") or ""
    template_rel = state.get("collection_template_path") or meta.get(
        "collection_template_path"
    )
    if not template_rel:
        return {"log_lines": ["跳过 collection_fill：无模板路径"]}

    template = root / template_rel
    if not template.is_file():
        return {
            "errors": [f"模板不存在: {template_rel}"],
            "log_lines": [f"collection_fill 失败：模板不存在"],
        }

    run_tool_script(
        root,
        "dump_sheet_headers.py",
        ["--src", template_rel, "--formulas"],
        timeout=60,
    )
    headers_path = root / "outputs" / "sheet_headers.json"
    headers_summary = ""
    if headers_path.is_file():
        headers_summary = headers_path.read_text(encoding="utf-8")[:8000]

    context = collect_ocr_snippets(root, max_files=14, max_total_chars=22000)
    llm = structured_llm(CollectionFillPlan)
    plan: CollectionFillPlan = llm.invoke(
        [
            ("system", prompts.COLLECTION_FILL_SYSTEM),
            (
                "human",
                prompts.COLLECTION_FILL_USER.format(
                    task_id=task_id,
                    template_path=template_rel,
                    headers_summary=headers_summary or "（无 headers 导出）",
                    context=context,
                ),
            ),
        ]
    )

    out_xlsx = root / "outputs" / "collection_filled.xlsx"
    shutil.copy2(template, out_xlsx)
    wb = load_workbook(out_xlsx)
    written = 0
    for cell in plan.cells:
        if cell.sheet not in wb.sheetnames:
            continue
        ws = wb[cell.sheet]
        col_idx = _col_to_index(cell.col)
        if col_idx is None or cell.row < 1:
            continue
        target = ws.cell(cell.row, col_idx)
        if isinstance(target, MergedCell):
            continue
        if target.data_type == "f" or (
            isinstance(target.value, str) and str(target.value).startswith("=")
        ):
            continue
        target.value = cell.value
        written += 1
    wb.save(out_xlsx)

    notes = root / "outputs" / "collection_gap_notes.txt"
    if plan.gap_notes:
        notes.write_text("\n".join(plan.gap_notes) + "\n", encoding="utf-8")

    return {
        "log_lines": [f"已生成 collection_filled.xlsx（写入 {written} 格）"],
        "progress": "信息收集表填报完成",
    }
