from __future__ import annotations

import json
import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


def _merge_anchor(ws, row: int, col: int) -> tuple[int, int]:
    """互斥组 D/E 常为纵向合并；非首行是 MergedCell，须写到合并区左上角。"""
    cell = ws.cell(row, col)
    if not isinstance(cell, MergedCell):
        return row, col
    for mr in ws.merged_cells.ranges:
        if mr.min_row <= row <= mr.max_row and mr.min_col <= col <= mr.max_col:
            return mr.min_row, mr.min_col
    return row, col


def _set_cell(ws, row: int, col: int, value) -> bool:
    anchor_row, anchor_col = _merge_anchor(ws, row, col)
    cell = ws.cell(anchor_row, anchor_col)
    if isinstance(cell, MergedCell):
        return False
    if cell.data_type == "f":
        return False
    if isinstance(cell.value, str) and cell.value.startswith("="):
        return False
    cell.value = value
    return True


def _exclusive_option_texts(root: Path) -> dict[tuple[str, int], str]:
    """互斥组每一行的选项原文，按 (sheet, row) 索引。

    D 列一律由程序按行号取原文写入，模型只负责指出选中哪一行。
    让模型复述选项全文会触发失控生成——2026-07-30 任务 19226215 的
    「周边配套分析」（选项 1 有 200+ 字带换行）连续两轮生成到 19.8 万字符
    仍未闭合 JSON。顺带也消除了「选项文本与 C 列不完全一致」这类校验失败。
    """
    catalog_path = root / "outputs" / "template_row_catalog.json"
    try:
        catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}

    texts: dict[tuple[str, int], str] = {}

    def walk(node, sheet: str = "") -> None:
        if isinstance(node, dict):
            sheet = str(node.get("name") or sheet)
            if node.get("choice_mode") == "exclusive":
                for row in node.get("rows") or []:
                    number = row.get("row")
                    option = row.get("option_text")
                    if isinstance(number, int) and option:
                        texts[(sheet, number)] = str(option)
            for value in node.values():
                walk(value, sheet)
        elif isinstance(node, list):
            for item in node:
                walk(item, sheet)

    walk(catalog)
    return texts


def write_xlsx_node(state: dict[str, Any]) -> dict[str, Any]:
    root = Path(state["extract_root"])
    template = root / "inputs" / "collection_template.xlsx"
    out = root / "outputs" / "collection_filled.xlsx"
    if not template.is_file():
        return {"errors": ["缺少 inputs/collection_template.xlsx"], "log_lines": ["write_xlsx 失败"]}

    shutil.copy2(template, out)
    option_texts = _exclusive_option_texts(root)
    fills = state.get("row_fills") or []
    # Prefer last fill for each (sheet, row) if duplicates from retries
    by_key: dict[tuple[str, int], dict] = {}
    for f in fills:
        rn = f.get("row")
        if rn is None:
            continue
        sheet = str(f.get("sheet") or "")
        by_key[(sheet, int(rn))] = f

    wb = load_workbook(out)
    default_sheet = wb.sheetnames[0]
    written_choice = 0
    written_remark = 0
    # 无 choice 的行先写备注；有 choice 的后写，避免互斥合并格被空选项备注覆盖
    ordered = sorted(
        by_key.items(),
        key=lambda kv: 0 if (kv[1].get("choice") in (None, "")) else 1,
    )
    for (sheet, row), f in ordered:
        ws_name = sheet if sheet in wb.sheetnames else default_sheet
        ws = wb[ws_name]
        choice = f.get("choice")
        remark = f.get("remark") or ""
        if choice not in (None, ""):
            # 互斥组不采信模型给的文本，一律按行号取模板原文
            value = option_texts.get((sheet, row), option_texts.get(("", row), choice))
            if _set_cell(ws, row, 4, value):  # D（合并则写到锚点）
                written_choice += 1
        if remark:
            if _set_cell(ws, row, 5, remark):  # E
                written_remark += 1
    wb.save(out)

    # notes markdown
    lines = [
        "| sheet | row | indicator_name | filled_choice | remark |",
        "|-------|-----|----------------|---------------|--------|",
    ]
    for (sheet, row) in sorted(by_key.keys(), key=lambda x: (x[0], x[1])):
        f = by_key[(sheet, row)]
        choice = f.get("choice") or ""
        remark = (f.get("remark") or "").replace("|", "\\|").replace("\n", " ")
        name = (f.get("indicator_name") or "").replace("|", "\\|")
        lines.append(f"| {sheet or default_sheet} | {row} | {name} | {choice} | {remark} |")
    notes_path = root / "outputs" / "collection_fill_notes.md"
    notes_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    return {
        "filled_path": "outputs/collection_filled.xlsx",
        "log_lines": [
            f"已写入 collection_filled.xlsx（D列={written_choice}, E列={written_remark}）"
        ],
        "progress": "指标表已写入 xlsx",
    }
