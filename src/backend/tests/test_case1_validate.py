from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook


def _add_scripts_to_path() -> None:
    scripts_dir = str((Path(__file__).resolve().parent.parent / "scripts"))
    if scripts_dir not in sys.path:
        sys.path.insert(0, scripts_dir)


_add_scripts_to_path()


# 真实任务里的正常备注，用作防误杀基准
_REAL_OK_REMARK = (
    "项目公司以其持有的成都市成华区迎晖路194号约33亩项目土地使用权提供第一顺位抵押担保。"
    "<引用>尽职调查报告.docx：原文「抵押担保：项目公司以其持有的位于成都市成华区迎晖路"
    "194号约33亩项目土地提供第一顺位抵押担保」</引用>"
)


def test_draft_reason_catches_thinking_markers():
    from case1_quality_rules import remark_draft_reason

    assert "等等，让我" in remark_draft_reason("楼面价高于阈值。等等，让我重新计算。")
    assert "再次检查材料" in remark_draft_reason("判断为否。\n\n再次检查材料：溢价率 0%")
    assert remark_draft_reason("先填否。修正判断：应为是") != ""


def test_draft_reason_catches_newlines_and_length():
    from case1_quality_rules import remark_draft_reason

    assert "换行" in remark_draft_reason("结论\n依据一\n依据二\n依据三")
    assert "长度" in remark_draft_reason("配套齐全。" * 200)


def test_draft_reason_does_not_flag_normal_remarks():
    from case1_quality_rules import remark_draft_reason

    assert remark_draft_reason(_REAL_OK_REMARK) == ""
    assert remark_draft_reason("") == ""
    # 「等等」作「诸如此类」用，不得误杀
    assert remark_draft_reason("周边有学校、医院、商场等等，配套齐全。") == ""
    # 单个换行属正常排版
    assert remark_draft_reason("结论。\n<引用>甲.docx：原文「乙」</引用>") == ""


def test_self_negation_only_matches_explicit_refusal():
    from case1_quality_rules import remark_self_negates

    assert remark_self_negates("周期不符，故不选此项。") == "故不选此项"
    assert remark_self_negates("因此不选此项") == "因此不选此项"
    assert remark_self_negates(_REAL_OK_REMARK) == ""
    # 选中项的正常理由不含自我否定
    assert remark_self_negates("满足出清周期与均价两项条件，故选此项。") == ""


def _write_case(
    tmp_path: Path,
    *,
    mode: str,
    choice: str,
    remark: str,
    explanation: str = "",
    option_texts: tuple[str, str] = ("1.出清周期快", "2.出清周期稳定"),
) -> tuple[Path, Path]:
    """造一个单指标组的 catalog + 已填 xlsx"""
    catalog = {
        "sheets": [
            {
                "name": "Sheet1",
                "sections": [
                    {
                        "indicator_groups": [
                            {
                                "indicator_name": "区域内一手房库存和去化情况",
                                "choice_mode": mode,
                                "indicator_explanation": explanation,
                                "allowed_choices": list(option_texts),
                                "rows": [
                                    {"row": 2, "option_text": option_texts[0]},
                                    {"row": 3, "option_text": option_texts[1]},
                                ],
                            }
                        ]
                    }
                ],
            }
        ]
    }
    catalog_path = tmp_path / "template_row_catalog.json"
    catalog_path.write_text(json.dumps(catalog, ensure_ascii=False), encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(2, 4).value = choice
    ws.cell(2, 5).value = remark
    # yes_no 每行都要填，否则会命中「未填指标选择」；exclusive 只能填一行
    if mode == "yes_no" and choice:
        ws.cell(3, 4).value = "否"
        ws.cell(3, 5).value = _REAL_OK_REMARK
    filled_path = tmp_path / "collection_filled.xlsx"
    wb.save(filled_path)
    return catalog_path, filled_path


def test_validate_reports_self_negation_with_indicator_name(tmp_path):
    """内容质量问题是 warning 不是 error，但须带「指标名」供修复轮定位重填组"""
    from validate_collection_filled import validate

    catalog_path, filled_path = _write_case(
        tmp_path,
        mode="exclusive",
        choice="2.出清周期稳定",
        remark="出清周期 4.28 个月，不满足区间，故不选此项。",
    )
    report = validate(catalog_path, filled_path)

    # 备注质量不该让整个任务失败——任务级 FAILED 只留给结构性缺失
    assert report["passed"] is True
    assert any("结论与选择相反" in w for w in report["warnings"])
    assert any("「区域内一手房库存和去化情况」" in w for w in report["warnings"])


def test_validate_reports_draft_remark_on_yes_no_row(tmp_path):
    from validate_collection_filled import validate

    catalog_path, filled_path = _write_case(
        tmp_path,
        mode="yes_no",
        choice="否",
        remark="楼面价低于阈值。等等，让我重新计算。12989*1.2=15586.8，所以应该是「是」。",
    )
    report = validate(catalog_path, filled_path)

    assert report["passed"] is True
    assert any("备注混入推演过程" in w for w in report["warnings"])


def test_quality_warnings_are_wired_to_refill_trigger():
    """降级成 warning 后仍须触发重填，否则等于把问题静音"""
    from app.agents.nodes.case1.validate_fix import _RETRY_WARNING_MARKERS

    samples = [
        "行33「某指标」备注称「故不选此项」却填了该选项，结论与选择相反",
        "行31「某指标」备注混入推演过程（含思考过程标记「重新计算」），只应写最终判断要点与引用",
        "行43「某指标」备注缺少源文件名与原文引用（建议：…）",
    ]
    for s in samples:
        assert any(m in s for m in _RETRY_WARNING_MARKERS), s
    # 豁免提示是给人看的，不该触发重填——模型判断不了项目属不属于豁免情形
    hint = "行13「某指标」填「否」，但该条目带豁免条款（…不适用按满分处理）——请确认"
    assert not any(m in hint for m in _RETRY_WARNING_MARKERS)


def test_judgement_thresholds_are_overridable(monkeypatch):
    """判据口径属内部契约，具体词表与阈值随客户模板而变，须可覆盖"""
    import importlib

    import case1_quality_rules as vcf

    monkeypatch.setenv("CASE1_DRAFT_MARKERS", "咱们再捋一遍,重新盘一下")
    monkeypatch.setenv("CASE1_REMARK_MAX_NEWLINES", "99")
    reloaded = importlib.reload(vcf)
    try:
        assert reloaded.remark_draft_reason("先填否。咱们再捋一遍。") != ""
        # 默认词表已被替换，原有标记不再命中
        assert reloaded.remark_draft_reason("让我重新计算一下") == ""
        # 换行阈值调高后，多换行不再判为草稿
        assert reloaded.remark_draft_reason("a\nb\nc\nd\ne") == ""
    finally:
        monkeypatch.undo()
        importlib.reload(vcf)


def test_exemption_clause_matches_template_wordings():
    from case1_quality_rules import exemption_clause, exemption_clauses

    # 当前模板里实际出现的两种写法
    assert exemption_clause("（如果为土地前融项目，该条目不适用，按满分处理）") != ""
    assert (
        exemption_clause("余额不少于5%（如果为介入前尚未开始预售的项目，该条目不适用按满分处理）")
        != ""
    )
    # 换客户/换模板后可能出现的其他措辞与半角括号
    assert exemption_clause("(土地前融项目视同满分)") != ""
    assert exemption_clause("（该项不计分）") != ""
    assert exemption_clause("", "考察增信措施落实情况") == ""
    assert exemption_clauses(
        "（土地前融项目不适用，按满分处理）另（未预售项目视同满分）"
    ) == ["（土地前融项目不适用，按满分处理）", "（未预售项目视同满分）"]


def test_exemption_markers_are_overridable(monkeypatch):
    import importlib

    import case1_quality_rules

    monkeypatch.setenv("CASE1_EXEMPTION_MARKERS", "豁免计分")
    reloaded = importlib.reload(case1_quality_rules)
    try:
        assert reloaded.exemption_clause("（该项目豁免计分）") != ""
        assert reloaded.exemption_clause("（该项目不适用，按满分处理）") == ""
    finally:
        monkeypatch.undo()
        importlib.reload(case1_quality_rules)


def test_validate_hints_exemption_on_negative_yes_no(tmp_path):
    from validate_collection_filled import validate

    clause = "（如果为介入前尚未开始预售的项目，该条目不适用按满分处理）"
    catalog_path, filled_path = _write_case(
        tmp_path,
        mode="yes_no",
        choice="否",
        remark=_REAL_OK_REMARK,
        option_texts=(f"□ 监管账户余额不少于融资本金5%{clause}", "□ 其他"),
    )
    report = validate(catalog_path, filled_path)

    assert report["passed"] is True  # 只提示，不拦
    assert any("带豁免条款" in w for w in report["warnings"])


def test_validate_hints_exemption_when_not_best_option(tmp_path):
    """选中项全文写在组首行 D 列，故须按文本而非行号判断是否选了最优项"""
    from validate_collection_filled import validate

    catalog_path, filled_path = _write_case(
        tmp_path,
        mode="exclusive",
        choice="2.常规监管",
        remark=_REAL_OK_REMARK,
        explanation="考察投后监管程度（如果为土地前融项目，该条目不适用，按满分处理）",
        option_texts=("1.强力监管", "2.常规监管"),
    )
    report = validate(catalog_path, filled_path)

    assert any("未选最优项" in w for w in report["warnings"])


def test_validate_no_exemption_hint_when_best_option_picked(tmp_path):
    from validate_collection_filled import validate

    catalog_path, filled_path = _write_case(
        tmp_path,
        mode="exclusive",
        choice="1.强力监管",
        remark=_REAL_OK_REMARK,
        explanation="考察投后监管程度（如果为土地前融项目，该条目不适用，按满分处理）",
        option_texts=("1.强力监管", "2.常规监管"),
    )
    report = validate(catalog_path, filled_path)

    assert not any("豁免条款" in w for w in report["warnings"])


def test_validate_detects_empty_yes_no_row_without_failing_task(tmp_path):
    """留空必须检出（否则成为重填时的逃逸路径），但默认不让整个任务失败

    材料千差万别，个别行确实可能无从判断；堵逃逸靠的是检出后触发重填，
    不靠 FAILED。需要严格模式的场景另有开关。
    """
    from validate_collection_filled import validate

    catalog_path, filled_path = _write_case(
        tmp_path, mode="yes_no", choice="", remark=""
    )
    report = validate(catalog_path, filled_path)

    assert report["passed"] is True
    assert any("未填指标选择" in w for w in report["warnings"])
    assert any("「区域内一手房库存和去化情况」" in w for w in report["warnings"])


def test_empty_choice_can_be_escalated_to_error(monkeypatch, tmp_path):
    import importlib

    import case1_quality_rules
    import validate_collection_filled

    monkeypatch.setenv("CASE1_EMPTY_CHOICE_LEVEL", "error")
    importlib.reload(case1_quality_rules)
    vcf = importlib.reload(validate_collection_filled)
    try:
        catalog_path, filled_path = _write_case(
            tmp_path, mode="yes_no", choice="", remark=""
        )
        report = vcf.validate(catalog_path, filled_path)
        assert report["passed"] is False
        assert any("未填指标选择" in e for e in report["errors"])
    finally:
        monkeypatch.undo()
        importlib.reload(case1_quality_rules)
        importlib.reload(validate_collection_filled)


def test_validate_passes_clean_fill(tmp_path):
    """正常填报不得被新判据打回"""
    from validate_collection_filled import validate

    catalog_path, filled_path = _write_case(
        tmp_path,
        mode="exclusive",
        choice="1.出清周期快",
        remark=_REAL_OK_REMARK,
    )
    report = validate(catalog_path, filled_path)

    assert report["passed"] is True
    assert report["errors"] == []
