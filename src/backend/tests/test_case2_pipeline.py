from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.agents import llm as llm_module
from app.agents.llm import structured_llm
from app.agents.nodes.case2.evidence import (
    _carryforward_review_flags,
    _entity_match_kind,
    _missing_ocr_sources,
    _route_evidence_chunks,
    _validate_sheet_identity,
    build_case2_ocr_chunks,
    facts_for_fill,
    facts_for_item_ids,
)
from app.agents.nodes.case2.fill_items import (
    _annotate_unmapped_columns,
    _normalize_value,
)
from app.agents.schemas.case2_item import (
    Case2BatchFill,
    Case2ChunkEvidence,
    Case2ItemFill,
    Case2PeriodMap,
)
from app.services.case2_amounts import is_repaired, parse_amount
from app.services.case2_defaults import MAX_FILL_LOGIC_RULES_LEN
from app.services.case2_schema_pipeline import (
    backfill_case2_schema,
    validate_case2_backfill,
)


def test_case2_ocr_chunks_cover_entire_markdown(tmp_path: Path):
    ocr_dir = tmp_path / "ocr_text"
    ocr_dir.mkdir()
    markers = [f"MARKER_{index:02d}" for index in range(12)]
    text = "\n".join(f"{marker} " + ("内容" * 45) for marker in markers)
    (ocr_dir / "report.md").write_text(text, encoding="utf-8")

    chunks = build_case2_ocr_chunks(tmp_path, max_chars=260, overlap=40)

    assert len(chunks) > 1
    combined = "\n".join(chunk["text"] for chunk in chunks)
    assert all(marker in combined for marker in markers)
    assert {chunk["source_ref"] for chunk in chunks} == {"ocr_text/report.md"}


def test_structured_llm_honors_explicit_method_for_openai(monkeypatch):
    class FakeModel:
        def __init__(self):
            self.calls = []

        def with_structured_output(self, schema, *, method):
            self.calls.append((schema, method))
            return "bound"

    model = FakeModel()
    monkeypatch.setattr(llm_module, "_is_bailian_family", lambda: False)

    result = structured_llm(dict, model=model, method="json_mode")

    assert result == "bound"
    assert model.calls == [(dict, "json_mode")]


def test_case2_ocr_chunks_fall_back_to_sidecar(tmp_path: Path):
    ocr_dir = tmp_path / "ocr_text"
    ocr_dir.mkdir()
    (ocr_dir / "scan.pdf.md").write_text("<!-- image -->", encoding="utf-8")
    (ocr_dir / "scan.pdf.ocr_cells.json").write_text(
        json.dumps(
            {
                "cells": [
                    {"page": 1, "text": "营业收入 100"},
                    {"page": 2, "text": "净利润 20"},
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    chunks = build_case2_ocr_chunks(tmp_path, max_chars=260, overlap=40)

    assert len(chunks) == 1
    assert "营业收入 100" in chunks[0]["text"]
    assert "净利润 20" in chunks[0]["text"]


def test_case2_ocr_chunks_preserve_scanned_page_markers(tmp_path: Path):
    ocr_dir = tmp_path / "ocr_text"
    ocr_dir.mkdir()
    (ocr_dir / "scan.md").write_text(
        "<!-- image -->\n第一页 " + ("甲" * 100)
        + "\n<!-- image -->\n第二页 "
        + ("乙" * 100),
        encoding="utf-8",
    )

    chunks = build_case2_ocr_chunks(tmp_path, max_chars=260, overlap=40)

    assert [chunk["page"] for chunk in chunks] == [1, 2]
    assert "第一页" in chunks[0]["text"]
    assert "第二页" in chunks[1]["text"]


def test_case2_ocr_chunks_restore_missing_page_header_from_sidecar(
    tmp_path: Path,
):
    ocr_dir = tmp_path / "ocr_text"
    ocr_dir.mkdir()
    (ocr_dir / "scan.pdf.md").write_text(
        "<!-- image -->\n第一页正文 "
        + ("甲" * 100)
        + "\n<!-- image -->\n| 营业收入 | 100 |",
        encoding="utf-8",
    )
    (ocr_dir / "scan.pdf.ocr_cells.json").write_text(
        json.dumps(
            {
                "cells": [
                    {
                        "page": 2,
                        "text": "利润表",
                        "bbox": [0.4, 0.05, 0.6, 0.1],
                    },
                    {
                        "page": 2,
                        "text": "2022年12期",
                        "bbox": [0.4, 0.1, 0.6, 0.15],
                    },
                    {
                        "page": 2,
                        "text": "营业收入 100",
                        "bbox": [0.1, 0.5, 0.9, 0.55],
                    },
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    chunks = build_case2_ocr_chunks(tmp_path, max_chars=500, overlap=40)

    assert chunks[1]["page"] == 2
    assert "OCR 页眉补充" in chunks[1]["text"]
    assert "利润表" in chunks[1]["text"]
    assert "2022年12期" in chunks[1]["text"]
    assert "营业收入 100" not in chunks[1]["text"]


def test_case2_fact_filter_and_value_normalization():
    catalog = {
        "facts": [
            {"item_id": "利润表:r5", "value": 10},
            {"item_id": "利润表:r6", "value": 20},
        ]
    }
    assert facts_for_item_ids(catalog, {"利润表:r6"}) == [
        {"item_id": "利润表:r6", "value": 20}
    ]
    assert _normalize_value("1,234.50", "number") == 1234.5
    assert _normalize_value("(12.5)", "number") == -12.5
    assert _normalize_value("null", "number") is None


def test_case2_schema_rejects_corrupted_structured_output():
    with pytest.raises(ValidationError):
        Case2ItemFill(
            item_id="利润表:r5",
            fields={"C": {"value": 1}},
            confidence="reason_one_line",
        )


def test_case2_item_unwraps_json_mode_field_values():
    item = Case2ItemFill.model_validate(
        {
            "item_id": "利润表:r6",
            "fields": {
                "C": {
                    "value": 100,
                    "evidence": "营业收入 100",
                }
            },
        }
    )

    assert item.fields == {"C": 100}


def test_case2_schema_wraps_single_objects_for_array_fields():
    evidence = Case2ChunkEvidence.model_validate(
        {
            "period_hints": {
                "statement_name": "利润表",
                "source_period": "本期金额",
                "evidence_text": "本期金额",
            },
            "facts": {
                "item_id": "利润表:r6",
                "subject_name": "营业收入",
                "source_period": "本期金额",
                "value": 100,
                "evidence_text": "营业收入 100",
            },
        }
    )
    batch = Case2BatchFill.model_validate(
        {"items": {"item_id": "利润表:r6", "fields": {"C": 100}}}
    )
    period_map = Case2PeriodMap.model_validate(
        {
            "columns": {
                "sheet_name": "利润表",
                "field_key": "C",
                "column_label": "最近一期",
            }
        }
    )

    assert len(evidence.period_hints) == 1
    assert len(evidence.facts) == 1
    assert evidence.facts[0].subject_name == "营业收入"
    assert len(batch.items) == 1
    assert len(period_map.columns) == 1


def test_case2_period_map_accepts_json_mode_aliases():
    period_map = Case2PeriodMap.model_validate(
        {
            "columns": {
                "template_column": "C",
                "report_date": "2023-03-31",
                "evidence": "最近一期报告",
            }
        }
    )

    assert period_map.columns[0].field_key == "C"
    assert period_map.columns[0].evidence_text == "最近一期报告"


def test_case2_schema_parses_labeled_period_hints():
    evidence = Case2ChunkEvidence.model_validate(
        {
            "period_hints": [
                "企业主体: 示例公司",
                "报告日期: 2023年12月31日",
            ],
            "facts": [],
        }
    )

    assert len(evidence.period_hints) == 1
    assert evidence.period_hints[0].entity_name == "示例公司"
    assert evidence.period_hints[0].report_date == "2023年12月31日"


def test_case2_fact_ignores_redundant_unknown_fields():
    evidence = Case2ChunkEvidence.model_validate(
        {
            "period_hints": [],
            "facts": {
                "item_id": "利润表:r6",
                "source_period": "本期金额",
                "value": 100,
                "evidence_text": "营业收入 100",
                "period_hints": ["报告日期: 2023年12月31日"],
            },
        }
    )

    assert evidence.facts[0].model_dump().get("period_hints") is None


def test_case2_period_hint_without_evidence_is_not_kept(monkeypatch):
    from app.agents.nodes.case2 import evidence as evidence_module

    class FakeRunnable:
        def invoke(self, _messages, **_kwargs):
            return Case2ChunkEvidence.model_validate(
                {
                    "period_hints": {
                        "statement_name": "利润表",
                        "source_period": "本期金额",
                    },
                    "facts": [],
                }
            )

    monkeypatch.setattr(
        evidence_module,
        "structured_llm",
        lambda *_args, **_kwargs: FakeRunnable(),
    )

    result = evidence_module.extract_evidence_chunk_node(
        {
            "task_id": "task",
            "material_chunk": {
                "source_ref": "ocr_text/report.md",
                "page": 1,
                "text": "利润表",
                "statement_hints": ["利润表"],
            },
            "evidence_targets": [],
            "chunk_index": 0,
            "total_chunks": 1,
        }
    )

    assert result["chunk_evidence"][0]["period_hints"] == []


def test_case2_null_fact_is_not_kept(monkeypatch):
    from app.agents.nodes.case2 import evidence as evidence_module

    class FakeRunnable:
        def invoke(self, _messages, **_kwargs):
            return Case2ChunkEvidence.model_validate(
                {
                    "period_hints": [],
                    "facts": {
                        "item_id": "利润表:r4",
                        "subject_name": "报告期",
                        "source_period": "本期金额",
                        "value": None,
                        "evidence_text": "2023年12期",
                    },
                }
            )

    monkeypatch.setattr(
        evidence_module,
        "structured_llm",
        lambda *_args, **_kwargs: FakeRunnable(),
    )

    result = evidence_module.extract_evidence_chunk_node(
        {
            "task_id": "task",
            "material_chunk": {
                "source_ref": "ocr_text/report.md",
                "page": 1,
                "text": "利润表",
            },
            "evidence_targets": [{"item_id": "利润表:r4"}],
            "chunk_index": 0,
            "total_chunks": 1,
        }
    )

    assert result["chunk_evidence"][0]["facts"] == []


def test_case2_routes_all_sources_but_only_financial_candidates():
    chunks = [
        {
            "source_ref": "ocr_text/audit.pdf.md",
            "source_name": "audit.pdf",
            "text": "合并资产负债表 货币资金 100",
        },
        {
            "source_ref": "ocr_text/audit.pdf.md",
            "source_name": "audit.pdf",
            "text": "审计意见和会计政策说明",
        },
        {
            "source_ref": "ocr_text/report.docx.md",
            "source_name": "report.docx",
            "text": "利润表分析 营业收入 200",
        },
    ]
    targets = [
        {
            "item_id": "资产负债表:r6",
            "sheet_name": "资产负债表",
            "label": "货币资金",
        },
        {
            "item_id": "利润表:r6",
            "sheet_name": "利润表",
            "label": "营业收入",
        },
    ]

    routed, inventory = _route_evidence_chunks(chunks, targets)

    assert len(routed) == 2
    assert {chunk["source_ref"] for chunk in routed} == {
        "ocr_text/audit.pdf.md",
        "ocr_text/report.docx.md",
    }
    assert sum(item["total_chunks"] for item in inventory) == 3
    assert sum(item["candidate_chunks"] for item in inventory) == 2


def test_case2_detects_source_without_ocr_text(tmp_path: Path):
    source_dir = tmp_path / "sources"
    source_dir.mkdir()
    (source_dir / "a.pdf").write_bytes(b"pdf")
    (source_dir / "b.docx").write_bytes(b"docx")

    missing = _missing_ocr_sources(
        tmp_path,
        [{"source_name": "a.pdf"}],
    )

    assert missing == ["b.docx"]


def _template_check_book(tmp_path: Path, formulas: dict[str, str], values: dict):
    """造一张带核查公式的小模板，行为与客户模板同构。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "资产负债表"
    for coord, value in values.items():
        ws[coord] = value
    for coord, formula in formulas.items():
        ws[coord] = formula
    path = tmp_path / "filled.xlsx"
    wb.save(path)
    return path


def test_template_checks_report_delta_and_locate_suspect_cell(tmp_path: Path):
    from app.services.case2_template_checks import (
        evaluate_checks,
        read_template_checks,
        summarize,
    )

    # 归母 + 少数股东权益 = 所有者权益合计，此处归母被漏填成 0
    path = _template_check_book(
        tmp_path,
        {
            "C50": '=IF(ABS(SUM(C10:C11)-C12)<1,"无误","所有者权益合计有误")',
            "C51": '=IF(ABS(SUM(C10:C11)-C12)<1,"无误","另一条")',
        },
        {"C10": 0, "C11": 46_418_122.14, "C12": 106_959_211.58},
    )
    outcomes = evaluate_checks(path, read_template_checks(path))
    assert summarize(outcomes)["failed"] == 2
    failed = outcomes[0]
    assert failed.check.message == "所有者权益合计有误"
    assert abs(failed.delta + 60_541_089.44) < 0.01


def test_template_checks_flag_cell_equal_to_delta(tmp_path: Path):
    """差额恰好等于某个被引用格的值时，要把那一格点出来。"""
    from app.services.case2_template_checks import evaluate_checks, read_template_checks

    # C11 被重复计入，差额正好等于 C11
    path = _template_check_book(
        tmp_path,
        {"C50": '=IF(ABS(SUM(C10:C11)-C12)<1,"无误","合计有误")'},
        {"C10": 100.0, "C11": 25.0, "C12": 100.0},
    )
    outcome = evaluate_checks(path, read_template_checks(path))[0]
    assert outcome.ok is False
    assert outcome.delta_matches == ["C11"]


def test_template_checks_solve_single_empty_cell(tmp_path: Path):
    from app.services.case2_template_checks import evaluate_checks, read_template_checks

    path = _template_check_book(
        tmp_path,
        {"C50": '=IF(ABS(SUM(C10:C11)-C12)<1,"无误","合计有误")'},
        {"C11": 46_418_122.14, "C12": 106_959_211.58},  # C10 真空着
    )
    outcome = evaluate_checks(path, read_template_checks(path))[0]
    assert outcome.empty_refs == ["C10"]
    assert outcome.suggestion is not None
    coord, value = outcome.suggestion
    assert coord == "C10"
    assert abs(value - 60_541_089.44) < 0.01


def test_template_checks_mark_unsupported_formula_as_unchecked(tmp_path: Path):
    """认不出的公式必须标「未校验」，绝不能当成通过。"""
    from app.services.case2_template_checks import (
        check_review_flags,
        evaluate_checks,
        read_template_checks,
        summarize,
    )

    path = _template_check_book(
        tmp_path,
        {
            "C50": '=IF(VLOOKUP(C10,A1:B9,2,FALSE)>0,"无误","不支持的写法")',
            "C51": "=C4",
        },
        {"C10": 1.0},
    )
    outcomes = evaluate_checks(path, read_template_checks(path))
    stats = summarize(outcomes)
    assert stats["passed"] == 0
    assert stats["unchecked"] == 1
    assert stats["mirror_cells"] == 1  # =C4 是镜像取值，不计入未校验
    assert any(f["kind"] == "template_check_unparsed" for f in check_review_flags(outcomes))


def test_case2_unmapped_column_is_excluded_from_filling():
    """未确定报告期的列不得参与填报，否则会填出不知属于哪一期的数值。"""
    from app.agents.nodes.case2.fill_items import _drop_unmapped_columns

    batches = [
        {
            "sheet_name": "利润表",
            "column_headers": {"C": "最近一期", "F": "上上期"},
            "items": [
                {
                    "item_id": "利润表:r6",
                    "fields": {
                        "C": {"cell": "C6", "value": None},
                        "F": {"cell": "F6", "value": None},
                    },
                }
            ],
        }
    ]
    period_mapping = {
        "columns": [
            {"sheet_name": "利润表", "field_key": "C", "report_date": "2023-12-31"},
            {"sheet_name": "利润表", "field_key": "F", "report_date": None},
        ]
    }

    dropped = _drop_unmapped_columns(batches, period_mapping)

    assert dropped == 1
    assert list(batches[0]["column_headers"]) == ["C"]
    assert list(batches[0]["items"][0]["fields"]) == ["C"]


def test_case2_period_matching_ignores_source_file_boundary():
    """同一报告期的数据可能散落在多份材料里，来源文件只作排序偏好。"""
    from app.agents.nodes.case2.evidence import _enrich_validated_column

    entries = [
        {
            "entity_name": "新鸿隆祥地产集团有限公司",
            "statement_scope": "合并",
            "statement_name": "利润表",
            "report_date": "2021-12-31",
            "source_ref": "ocr_text/sources/2021.pdf.md",
            "evidence_text": "本期数",
        }
    ]
    # 列指向 2022 年报告（比较列），事实却在 2021 年报告里
    column = {
        "sheet_name": "利润表",
        "field_key": "E",
        "entity_name": "新鸿隆祥地产集团有限公司",
        "statement_scope": "合并",
        "statement_name": "利润表",
        "report_date": "2021-12-31",
        "source_ref": "ocr_text/sources/2022.pdf.md",
    }

    out = _enrich_validated_column(column, sheet_name="利润表", entries=entries)

    assert out["report_date"] == "2021-12-31"


def test_case2_period_date_recovered_from_evidence_text():
    """模型把推断出的日期只写进说明文本时取回；只有年份则不取。"""
    from app.agents.nodes.case2.evidence import _recover_report_date

    mapped = {
        "report_date": None,
        "evidence_text": "2022年审计报告利润表显示'上年累计金额'列，对应上一年度2021-12-31。",
    }
    assert _recover_report_date(mapped) is True
    assert mapped["report_date"] == "2021-12-31"

    vague = {
        "report_date": None,
        "evidence_text": "源材料中未提供2020年审计报告，根据序列推断为2020年，但无直接证据。",
    }
    assert _recover_report_date(vague) is False
    assert vague["report_date"] is None

    already = {"report_date": "2023-12-31", "evidence_text": "对应最新报告期2023-12-31。"}
    assert _recover_report_date(already) is False


def test_case2_period_map_accepts_unmappable_column_with_blank_fields():
    """材料缺某一期时模型会把该列各字段留空，不能让整份映射解析失败。"""
    period_map = Case2PeriodMap.model_validate(
        {
            "columns": [
                {
                    "sheet_name": "利润表",
                    "field_key": "C",
                    "statement_scope": "合并",
                    "report_date": "2023-12-31",
                },
                {
                    "sheet_name": "利润表",
                    "field_key": "F",
                    "statement_scope": "",
                    "report_date": "",
                    "entity_name": "",
                    "source_ref": "",
                    "confidence": "low",
                    "evidence_text": "材料中无 2020 年完整年报，无法映射",
                },
            ]
        }
    )

    unmapped = period_map.columns[1]
    assert unmapped.statement_scope == "未知"
    assert unmapped.report_date is None
    assert unmapped.entity_name is None


def test_upload_zip_md5_stable_for_same_sources(tmp_path: Path):
    """相同源文件须打出相同 MD5，历史任务的 OCR 结果才能被复用。"""
    import hashlib

    from app.services.upload_zip_builder import build_zip_from_pairs

    pairs = [("b.pdf", b"world" * 100), ("a.pdf", b"hello" * 100)]

    def md5_of(name: str, extras) -> str:
        dest = tmp_path / name
        build_zip_from_pairs(dest, extras=extras, extras_prefix="sources")
        return hashlib.md5(dest.read_bytes()).hexdigest()

    first = md5_of("one.zip", pairs)
    assert first == md5_of("two.zip", pairs)
    assert first == md5_of("three.zip", list(reversed(pairs)))
    assert first != md5_of(
        "four.zip", [("b.pdf", b"world" * 100), ("a.pdf", b"HELLO" * 100)]
    )


def _calc_module():
    import sys
    from pathlib import Path as _P

    scripts = str(_P(__file__).resolve().parent.parent / "scripts")
    if scripts not in sys.path:
        sys.path.insert(0, scripts)
    import apply_case2_calc_rules

    return apply_case2_calc_rules


def test_calc_sum_overwrites_model_value_and_flags_it():
    """sum/diff 以规则为准，覆盖已有值并留痕。"""
    calc = _calc_module()
    data = {"sheets": [{"sheet": "资产负债表", "items": [
        {"item_id": "r9", "label": "应收票据及应收账款",
         "fields": {"D": {"cell": "D9", "value": 465147735.17}}},
        {"item_id": "r10", "label": "其中:应收票据",
         "fields": {"D": {"cell": "D10", "value": 110516.93}}},
        {"item_id": "r11", "label": "应收账款",
         "fields": {"D": {"cell": "D11", "value": 354537218.24}}},
    ]}]}
    report = calc.apply_calc_rules(data, [{
        "op": "sum", "sheet": "资产负债表",
        "target_label": "应收票据及应收账款",
        "source_labels": ["其中应收票据", "应收账款"],
    }])

    assert data["sheets"][0]["items"][0]["fields"]["D"]["value"] == 354647735.17
    assert report["applied"][0]["overwrote"] == 465147735.17
    assert any(f["kind"] == "calc_overwrote_model_value" for f in report["review_flags"])


def test_calc_sum_partial_sources_fill_empty_but_never_overwrite():
    """来源缺值时部分和只能补空格，不得覆盖模型已填的完整值。"""
    calc = _calc_module()
    data = {"sheets": [{"sheet": "资产负债表", "items": [
        {"item_id": "r64", "label": "应付票据及应付账款",
         "fields": {"D": {"cell": "D64", "value": 425997938.46},
                    "C": {"cell": "C64", "value": None}}},
        {"item_id": "r65", "label": "其中:应付票据",
         "fields": {"D": {"cell": "D65", "value": None},
                    "C": {"cell": "C65", "value": None}}},
        {"item_id": "r66", "label": "应付账款",
         "fields": {"D": {"cell": "D66", "value": 325897828.35},
                    "C": {"cell": "C66", "value": 342192538.05}}},
    ]}]}
    report = calc.apply_calc_rules(data, [{
        "op": "sum", "sheet": "资产负债表",
        "target_label": "应付票据及应付账款",
        "source_labels": ["其中应付票据", "应付账款"],
    }])

    fields = data["sheets"][0]["items"][0]["fields"]
    # D 已有值且来源不全：保留模型值
    assert fields["D"]["value"] == 425997938.46
    assert any("incomplete sources" in str(s.get("reason")) for s in report["skipped"])
    # C 为空：部分和可以补
    assert fields["C"]["value"] == 342192538.05


def test_calc_diff_subtracts_and_requires_all_sources():
    calc = _calc_module()
    data = {"sheets": [{"sheet": "资产负债表", "items": [
        {"item_id": "r119", "label": "归属于母公司所有者权益合计",
         "fields": {"F": {"cell": "F119", "value": None}, "E": {"cell": "E119", "value": None}}},
        {"item_id": "r120", "label": "少数股东权益",
         "fields": {"F": {"cell": "F120", "value": 46418122.14}, "E": {"cell": "E120", "value": None}}},
        {"item_id": "r121", "label": "所有者权益合计",
         "fields": {"F": {"cell": "F121", "value": 106959211.58}, "E": {"cell": "E121", "value": 601800680.28}}},
    ]}]}
    report = calc.apply_calc_rules(data, [{
        "op": "diff", "sheet": "资产负债表",
        "target_label": "归属于母公司所有者权益合计",
        "source_labels": ["所有者权益合计", "少数股东权益"],
    }])

    fields = data["sheets"][0]["items"][0]["fields"]
    assert fields["F"]["value"] == 60541089.44
    # E 列少数股东权益缺值：差额缺任一来源都不算，不出半截差额
    assert fields["E"]["value"] is None
    assert any("missing source value" in str(s.get("reason")) for s in report["skipped"])


def test_llm_guard_aborts_whitespace_degeneration_in_streaming():
    """量化模型在 JSON 冒号后无限吐空格时，必须在流式过程中被中止。

    这里走真实的 langchain 流式管道，因为 handler 抛的异常默认会被
    handle_event 吞掉，只有 raise_error=True 才会向上传播。
    """
    from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
    from langchain_core.messages import AIMessage

    from app.agents.llm import LLMDegenerationError, guard_config

    degenerate = '{"confidence":' + " " * 3000 + "1}"
    model = GenericFakeChatModel(messages=iter([AIMessage(content=degenerate)]))
    with pytest.raises(LLMDegenerationError):
        for _ in model.stream("hi", config=guard_config()):
            pass


def test_llm_guard_does_not_trip_on_normal_indented_json():
    from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
    from langchain_core.messages import AIMessage

    from app.agents.llm import guard_config

    pretty = json.dumps(
        {"facts": [{"item_id": f"r{i}", "value": i} for i in range(40)]},
        ensure_ascii=False,
        indent=4,
    )
    model = GenericFakeChatModel(messages=iter([AIMessage(content=pretty)]))
    chunks = list(model.stream("hi", config=guard_config()))
    assert chunks


def test_llm_guard_uses_fresh_handler_per_call():
    """证据抽取是 8 路并发，共用 handler 会把各路的空白游程算到一起。"""
    from app.agents.llm import guard_config

    first = guard_config()["callbacks"][0]
    second = guard_config()["callbacks"][0]
    assert first is not second


def test_llm_output_cap_applies_in_streaming_too(monkeypatch):
    """流式也必须设输出上限——这是唯一确定性的终止条件。

    2026-07-31 实测：读超时测的是字节间隔（正常响应最大间隔 0.41 秒，够不着
    180 秒），空白探测器只数连续空白（重复非空白内容不触发），当时流式分支
    直接返回 {}，于是证据抽取挂死 18 分钟无人叫停。
    """
    from app.agents import llm as llm_module

    monkeypatch.setattr(llm_module.settings, "llm_streaming", True, raising=False)
    assert llm_module.output_cap(32768) == {"max_tokens": 32768}

    monkeypatch.setattr(llm_module.settings, "llm_streaming", False, raising=False)
    assert llm_module.output_cap(32768) == {"max_tokens": 32768}


def test_case2_evidence_extraction_has_an_output_cap():
    """证据抽取曾是唯一没有上限的调用点，回归时必须守住。"""
    from app.agents.nodes.case2 import evidence as ev

    assert ev._EVIDENCE_MAX_TOKENS >= 32768


def test_case2_entity_ocr_typo_is_same_subject_but_group_suffix_is_not():
    """扫描件公司名错一个字要当同一家；「集团」这种成分差异不能当同一家。"""
    assert _entity_match_kind("东厦建设开发集团有限公司", "东度建设开发集团有限公司") == "fuzzy"
    assert (
        _entity_match_kind("东厦建设开发集团有限公司", "东厦建设开发团有限公司（并）")
        == "fuzzy"
    )
    assert _entity_match_kind("东厦建设开发集团有限公司", "东厦建设开发有限公司") == "none"
    assert _entity_match_kind("山东尊创置业有限公司", "山东新鸿置业有限公司") == "none"
    # 短名字差一个字不能当同一家
    assert _entity_match_kind("甲公司", "乙公司") == "none"


def test_case2_facts_survive_entity_name_ocr_typo():
    """同一份材料里公司名被 OCR 认错，不能因此丢掉整页证据。"""
    catalog = {
        "source_inventory": [{"source_ref": "ocr_text/a.pdf.md"}],
        "facts": [
            {
                "item_id": "资产负债表:r6",
                "entity_name": "东度建设开发集团有限公司",
                "statement_scope": "合并",
                "statement_name": "资产负债表",
                "report_date": "2022-12-31",
                "value": 248573298.6,
                "source_ref": "ocr_text/a.pdf.md",
            }
        ],
    }
    period_mapping = {
        "columns": [
            {
                "sheet_name": "资产负债表",
                "field_key": "D",
                "entity_name": "东厦建设开发集团有限公司",
                "statement_scope": "合并",
                "report_date": "2022-12-31",
            }
        ]
    }

    facts = facts_for_fill(
        catalog,
        {"资产负债表:r6"},
        sheet_name="资产负债表",
        period_mapping=period_mapping,
    )

    assert [fact["value"] for fact in facts] == [248573298.6]


def test_case2_carryforward_mismatch_is_flagged_without_changing_data():
    """本期年初数与上期期末数对不上时只报复核提示。"""
    def fact(report_date: str, period: str, subject: str, value: float, src: str) -> dict:
        return {
            "statement_name": "资产负债表",
            "report_date": report_date,
            "source_period": period,
            "subject_name": subject,
            "value": value,
            "source_ref": src,
        }

    # 来源必须给全：同一份报告的年初列与期末列本就是两期，不能互比，
    # 单份材料无法自证期末/年初是否颠倒，必须靠另一份报告交叉。
    subjects = ("货币资金", "应收账款", "资产总计", "流动资产合计")
    catalog = {
        "facts": [
            # 2020 年报的期末数
            *(fact("2020-12-31", "期末数", name, 100.0, "2020.md") for name in subjects),
            # 2021 年报的「年初数」应等于上面这组，但这里整体对不上
            *(fact("2021-12-31", "年初数", name, 900.0, "2021.md") for name in subjects),
            # 2021 年报的期末数，与 2022 年报的年初数一致，不应报
            *(fact("2021-12-31", "期末数", name, 500.0, "2021.md") for name in subjects),
            *(fact("2022-12-31", "年初数", name, 500.0, "2022.md") for name in subjects),
        ]
    }

    flags = _carryforward_review_flags(catalog)

    assert [flag["report_date"] for flag in flags] == ["2021-12-31"]
    assert flags[0]["kind"] == "carryforward_mismatch"
    # 只提示，不改动任何事实
    assert catalog["facts"][4]["value"] == 900.0


def test_case2_unmapped_column_reason_is_not_reported_as_missing_source():
    schema = {
        "sheets": [
            {
                "sheet": "资产负债表",
                "items": [
                    {
                        "item_id": "资产负债表:r6",
                        "fields": {
                            "C": {"cell": "C6", "value": None},
                            "E": {"cell": "E6", "value": None},
                        },
                        "reason_one_line": "文件中未发现相关信息",
                    },
                    {
                        "item_id": "资产负债表:r9",
                        "fields": {"E": {"cell": "E9", "value": 1.0}},
                        "reason_one_line": "文件中未发现相关信息",
                    },
                ],
            }
        ]
    }
    period_mapping = {
        "columns": [
            {
                "sheet_name": "资产负债表",
                "field_key": "C",
                "column_label": "最近一期报告",
                "report_date": None,
            },
            {
                "sheet_name": "资产负债表",
                "field_key": "E",
                "column_label": "上期(年报)",
                "report_date": "2021-12-31",
            },
        ]
    }

    _annotate_unmapped_columns(schema, period_mapping)
    items = schema["sheets"][0]["items"]

    assert "未映射成功" in items[0]["reason_one_line"]
    assert "非源文件缺失" in items[0]["reason_one_line"]
    # 有值的行不加注
    assert items[1]["reason_one_line"] == "文件中未发现相关信息"

    # 重复执行不叠加
    _annotate_unmapped_columns(schema, period_mapping)
    assert items[0]["reason_one_line"].count("未映射成功") == 1


def test_case2_facts_keep_one_subject_but_never_mix_statement_scopes():
    """一个任务默认一家公司，主体名不再过滤；合并与母公司仍不得混填。"""
    catalog = {
        "facts": [
            {
                "item_id": "利润表:r6",
                "entity_name": "甲公司",
                "statement_scope": "合并",
                "statement_name": "利润表",
                "report_date": "2023-12-31",
                "value": 100,
                "source_ref": "ocr_text/a.pdf.md",
            },
            {
                "item_id": "利润表:r6",
                "entity_name": "乙公司",
                "statement_scope": "合并",
                "statement_name": "利润表",
                "report_date": "2023-12-31",
                "value": 200,
                "source_ref": "ocr_text/b.pdf.md",
            },
            {
                "item_id": "利润表:r6",
                "entity_name": "甲公司",
                "statement_scope": "母公司",
                "statement_name": "利润表",
                "report_date": "2023-12-31",
                "value": 300,
                "source_ref": "ocr_text/c.pdf.md",
            },
        ]
    }
    period_mapping = {
        "columns": [
            {
                "sheet_name": "利润表",
                "field_key": "C",
                "entity_name": "甲公司",
                "statement_scope": "合并",
                "report_date": "2023-12-31",
                "source_ref": "ocr_text/a.pdf.md",
            }
        ]
    }

    facts = facts_for_fill(
        catalog,
        {"利润表:r6"},
        sheet_name="利润表",
        period_mapping=period_mapping,
    )

    # 主体名写法不同的证据照常参与；母公司口径的那条仍被排除
    assert sorted(fact["value"] for fact in facts) == [100, 200]


def test_case2_facts_match_entity_alias_and_date_format():
    catalog = {
        "facts": [
            {
                "item_id": "利润表:r6",
                "entity_name": "东厦",
                "statement_scope": "合并",
                "statement_name": "利润表",
                "report_date": "2022年12月31日",
                "value": 100,
                "source_ref": "ocr_text/a.pdf.md",
            }
        ]
    }
    period_mapping = {
        "columns": [
            {
                "sheet_name": "利润表",
                "field_key": "D",
                "entity_name": "东厦建设开发集团有限公司",
                "statement_scope": "合并",
                "report_date": "2022-12-31",
                "source_ref": "ocr_text/a.pdf.md",
            }
        ]
    }

    facts = facts_for_fill(
        catalog,
        {"利润表:r6"},
        sheet_name="利润表",
        period_mapping=period_mapping,
    )

    assert [fact["value"] for fact in facts] == [100]


def test_case2_period_mapping_does_not_abort_on_differing_entity_names():
    """一个任务默认一家公司：列上主体名不同不再终止任务，交复核提示处理。"""
    columns = [
        {"sheet_name": "利润表", "entity_name": "甲公司"},
        {"sheet_name": "利润表", "entity_name": "乙公司"},
    ]

    _validate_sheet_identity(columns)


def test_case2_period_mapping_still_rejects_mixed_statement_scopes():
    with pytest.raises(RuntimeError, match="多个报表口径"):
        _validate_sheet_identity(
            [
                {"sheet_name": "利润表", "statement_scope": "合并"},
                {"sheet_name": "利润表", "statement_scope": "母公司"},
            ]
        )


def test_case2_period_mapping_unifies_entity_ocr_variants():
    """同一主体被 OCR 认成几个变体时不能判为「混入多个企业主体」而终止任务。"""
    columns = [
        {"sheet_name": "资产负债表", "field_key": "C", "entity_name": "东厦建设开发团有限公司（并）"},
        {"sheet_name": "资产负债表", "field_key": "D", "entity_name": "东度建设开发集团有限公司"},
        {"sheet_name": "资产负债表", "field_key": "E", "entity_name": "东厦建设开发集团有限公司"},
    ]

    _validate_sheet_identity(columns, preferred_entity="东厦建设开发集团有限公司")

    # 统一到出现频次最高的写法，而不是被认错的那个
    assert {column["entity_name"] for column in columns} == {"东厦建设开发集团有限公司"}


def test_case2_period_mapping_normalizes_entity_abbreviations():
    columns = [
        {"sheet_name": "利润表", "entity_name": "东厦"},
        {
            "sheet_name": "利润表",
            "entity_name": "东厦建设开发集团有限公司",
        },
    ]

    _validate_sheet_identity(columns)

    assert {column["entity_name"] for column in columns} == {
        "东厦建设开发集团有限公司"
    }


def test_case2_validation_rejects_all_empty(tmp_path: Path):
    outputs = tmp_path / "outputs"
    outputs.mkdir()
    (outputs / "backfill_report.json").write_text(
        json.dumps({"written_cells": 0}), encoding="utf-8"
    )
    (outputs / "case2_filled_schema.json").write_text(
        json.dumps(
            {
                "sheets": [
                    {
                        "sheet": "利润表",
                        "items": [
                            {
                                "item_id": "利润表:r5",
                                "fields": {
                                    "C": {
                                        "cell": "C5",
                                        "value_type": "number",
                                        "value": None,
                                    }
                                },
                            }
                        ],
                    }
                ]
            }
        ),
        encoding="utf-8",
    )

    ok, message = validate_case2_backfill(tmp_path)

    assert ok is False
    assert "全空" in message


def test_case2_backfill_counts_only_nonempty_and_writes_real_types(tmp_path: Path):
    inputs = tmp_path / "inputs"
    outputs = tmp_path / "outputs"
    inputs.mkdir()
    outputs.mkdir()

    wb = Workbook()
    ws = wb.active
    ws.title = "利润表"
    wb.save(inputs / "collection_template.xlsx")
    wb.close()

    filled = {
        "sheets": [
            {
                "sheet": "利润表",
                "items": [
                    {
                        "item_id": "利润表:r4",
                        "evidence_refs": ["ocr_text/report.md"],
                        "fields": {
                            "C": {
                                "cell": "C4",
                                "value_type": "date",
                                "value": "2023-03-31",
                            }
                        },
                    },
                    {
                        "item_id": "利润表:r5",
                        "evidence_refs": ["ocr_text/report.md"],
                        "fields": {
                            "C": {
                                "cell": "C5",
                                "value_type": "number",
                                "value": "1,234.50",
                            },
                            "D": {
                                "cell": "D5",
                                "value_type": "number",
                                "value": None,
                            },
                        },
                    },
                ],
            }
        ]
    }
    (outputs / "case2_filled_schema.json").write_text(
        json.dumps(filled, ensure_ascii=False), encoding="utf-8"
    )

    ok, message = backfill_case2_schema(tmp_path)

    assert ok, message
    report = json.loads((outputs / "backfill_report.json").read_text())
    assert report["target_cells"] == 3
    assert report["written_cells"] == 2
    assert report["empty_cells"] == 1

    check = load_workbook(outputs / "collection_filled.xlsx", data_only=False)
    try:
        assert isinstance(check["利润表"]["C4"].value, datetime)
        assert check["利润表"]["C5"].value == 1234.5
        assert check["利润表"]["D5"].value is None
    finally:
        check.close()


def test_amount_separator_repair_keeps_normal_values_untouched():
    # 常规写法必须原样通过，重组只在常规解析失败时兜底
    assert parse_amount("457,265,859.05") == 457265859.05
    assert parse_amount("0.00") == 0.0
    assert parse_amount("1.5") == 1.5
    assert parse_amount("(1,234.56)") == -1234.56
    assert is_repaired("457,265,859.05") is False


def test_amount_separator_repair_recovers_ocr_confused_separators():
    # 均取自新鸿 2023 强制整页 OCR 的真实产物
    assert parse_amount("450,070.190.00") == 450070190.00
    assert parse_amount("1.758,823.768.78") == 1758823768.78
    assert parse_amount("8.785.378,359.75") == 8785378359.75
    assert parse_amount("10.704.829,000.89") == 10704829000.89
    assert is_repaired("1.758,823.768.78") is True


def test_amount_repair_refuses_ambiguous_or_garbled_text():
    # 分组不合千分位就不猜，宁可丢事实也不编数
    assert parse_amount("1.2.3") is None
    assert parse_amount("767、 7(Xl^00000") is None
    assert parse_amount("abc") is None
    assert parse_amount("") is None


def test_normalize_value_uses_amount_repair():
    assert _normalize_value("1.758,823.768.78", "number") == 1758823768.78
    assert _normalize_value("1,234.50", "number") == 1234.5


def _add_scripts_to_path() -> None:
    import sys

    scripts_dir = str((Path(__file__).resolve().parent.parent / "scripts"))
    if scripts_dir not in sys.path:
        sys.path.insert(0, scripts_dir)


_add_scripts_to_path()


def test_text_layer_guard_flags_only_garbled_amount_pages():
    from text_layer_guard import _is_wellformed

    assert _is_wellformed("399,093,581.74") is True
    assert _is_wellformed("0.00") is True
    assert _is_wellformed("399.093、 581 74") is False
    assert _is_wellformed("767、 7(Xl^00000") is False


def test_text_layer_guard_recognizes_genuinely_broken_amounts():
    """全部取自新鸿 2023 文本层的真实内容。"""
    from text_layer_guard import _is_broken_amount

    for text in (
        "767、 7(Xl^00000",
        "399.093、 581 74",
        "800`955.434 17",
        "457,265,85905",
        "450.070 19000",
    ):
        assert _is_broken_amount(text) is True, text


def test_text_layer_guard_does_not_flag_list_markers_or_dates():
    """回归：初版判据把个人征信报告的列表序号当成畸形金额，误判整份文件。

    误判方向是最坏的——会把一份准确的文本层扔掉换成 OCR 结果。
    """
    from text_layer_guard import _is_broken_amount

    for text in (
        "1.", "2.", "9.", "49",           # 征信报告的列表序号与条目数
        "2023-05-01", "2023.05.01",       # 日期
        "399,093,581.74", "-304,491,654.65", "0.00", "11,390,265,552.64",
    ):
        assert _is_broken_amount(text) is False, text


def test_ocr_guards_scoped_to_case2_by_default():
    """OCR 增强默认只对 case2 生效；case0 材料杂、误判代价高，先不启用。"""
    from app.services.ocr_runner import ocr_guards_enabled

    assert ocr_guards_enabled("case2") is True
    assert ocr_guards_enabled("case0") is False
    assert ocr_guards_enabled("case1") is False


def test_user_rules_prompt_limit_matches_task_create_limit():
    # 两处口径必须相同，否则超出部分会在进模型前被静默截掉
    from app.services import task_create

    assert task_create.MAX_FILL_LOGIC_RULES_LEN == MAX_FILL_LOGIC_RULES_LEN


def test_map_periods_prompt_separates_negotiable_rules():
    from app.agents.prompts.case2 import MAP_PERIODS_SYSTEM

    assert "【不可协商】" in MAP_PERIODS_SYSTEM
    assert "【默认推定】" in MAP_PERIODS_SYSTEM
    # 主体/口径隔离与禁编造必须留在不可协商段
    head, _, tail = MAP_PERIODS_SYSTEM.partition("【默认推定】")
    assert "同一企业主体和同一报表口径" in head
    assert "不得生成" in head
    assert "最近一期报告" in tail


_DOUBLE_COLUMN_TABLE = """## 资产负债表 2021-12-31

| 资产     | 行次 |                | 年初数 负债和所有者（或股东）权益 | 行次 | 期末数        | 年初数        |
|----------|------|----------------|--------------------------|------|---------------|---------------|
| 流动资产：|      |                | 流动负债：                | |               |               |
| 货币资金 | 1    | 310,893,340.43 | 134,715,961.79 短期借款   | 31   | 20,000,000.00 | 40,000,000.00 |
| 短期投资 | 2    | 0.00           | 0.00 应付票据             | 32   | 67,113,994.45 | 0.00          |
| 应收票据 | 3    | 699,000.00     | 3,700,000.00 应付账款     | 33   | 381,509,297.36| 1,297,816.00  |
| 应收账款 | 4    | 311,266,236.53 | 125,100,383.26 预收账款   | 34   | 715,641,831.00| 91,232.00     |
| 预付账款 | 5    | 20,634,073.53  | 15,217,548.83 应付职工薪酬 | 35   | 1,468,025.50  | 408,373.29    |
"""


def _split_rows(markdown: str) -> list[list[str]]:
    return [
        [cell.strip() for cell in line.strip().strip("|").split("|")]
        for line in markdown.splitlines()
        if line.startswith("|") and set(line) - set("|- ")
    ]


def test_double_column_table_is_split_back_into_two_halves():
    from table_split import split_double_column_tables

    out, reports = split_double_column_tables(_DOUBLE_COLUMN_TABLE)
    rows = _split_rows(out)

    assert len(reports) == 1
    assert reports[0]["columns_after"] == 8
    assert reports[0]["unsplittable_cells"] == []
    # 表头拆成「年初数」和负债侧科目列
    assert rows[0][3] == "年初数"
    assert rows[0][4].startswith("负债和所有者")
    # 数据行：左半年初数与右半科目名各归各位
    货币资金 = rows[2]
    assert 货币资金[2] == "310,893,340.43"
    assert 货币资金[3] == "134,715,961.79"
    assert 货币资金[4] == "短期借款"
    assert 货币资金[6] == "20,000,000.00"


def test_double_column_split_handles_interleaved_header_labels():
    # OCR 有时把两个表头格交错读成「负债和所有者 年初数 (或股东)权益」
    from table_split import split_double_column_tables

    markdown = _DOUBLE_COLUMN_TABLE.replace(
        "年初数 负债和所有者（或股东）权益", "负债和所有者 年初数 (或股东)权益"
    )
    out, reports = split_double_column_tables(markdown)
    rows = _split_rows(out)

    assert len(reports) == 1
    assert rows[0][3] == "年初数"
    assert rows[0][4] == "负债和所有者(或股东)权益"


def test_double_column_split_refuses_cells_it_cannot_resolve():
    # 一格里两个金额说明还发生了跨行错位，不猜：金额侧留空并记入 failures
    from table_split import split_double_column_tables

    markdown = _DOUBLE_COLUMN_TABLE.replace(
        "3,700,000.00 应付账款", "3,700,000.00 1,234.00 应付账款"
    )
    out, reports = split_double_column_tables(markdown)
    rows = _split_rows(out)

    assert reports[0]["unsplittable_cells"] == ["3,700,000.00 1,234.00 应付账款"]
    应收票据 = rows[4]
    assert 应收票据[3] == ""
    assert 应收票据[4] == "3,700,000.00 1,234.00 应付账款"


def test_double_column_split_leaves_ordinary_tables_untouched():
    from table_split import split_double_column_tables

    markdown = """| 项目     | 行次 | 本年累计金额     | 本月金额       |
|----------|------|------------------|----------------|
| 营业收入 | 1    | 1,114,184,729.27 | 260,731,541.20 |
| 营业成本 | 2    | 1,094,459,146.79 | 260,128,870.11 |
"""
    out, reports = split_double_column_tables(markdown)

    assert reports == []
    assert out == markdown


def test_reuse_requires_every_source_to_have_ocr_markdown(tmp_path: Path):
    """OCR 中途失败留下的残缺产物不能被当成「已 OCR」而跳过重跑。

    2026-07-31：新鸿三文件任务在第 3 份 OOM，目录里只剩 2021/2022 两份 md，
    同材料的新任务会命中它并整个跳过 OCR，缺的 2023 再也补不上。
    """
    from app.services.reuse_extract import extract_has_ocr_markdown

    src = tmp_path / "sources"
    src.mkdir()
    for name in ("a.pdf", "b.pdf", "c.pdf"):
        (src / name).write_bytes(b"%PDF-1.4\n")
    ocr = tmp_path / "ocr_text" / "sources"
    ocr.mkdir(parents=True)

    (ocr / "a.pdf.md").write_text("x", encoding="utf-8")
    assert extract_has_ocr_markdown(tmp_path) is False   # 残缺

    (ocr / "b.pdf.md").write_text("x", encoding="utf-8")
    assert extract_has_ocr_markdown(tmp_path) is False   # 仍残缺

    (ocr / "c.pdf.md").write_text("x", encoding="utf-8")
    assert extract_has_ocr_markdown(tmp_path) is True    # 齐了才算数


def test_case1_exclusive_choice_uses_template_text_not_model_output(tmp_path: Path):
    """互斥组 D 列一律按行号取模板原文，不采信模型复述的文本。

    2026-07-30 任务 19226215：「周边配套分析」选项 1 有 200+ 字带换行，
    模型被要求复述全文，连续两轮生成到 19.8 万字符仍未闭合 JSON，任务 FAILED。
    """
    import json

    from openpyxl import Workbook

    from app.agents.nodes.case1.write_xlsx import _exclusive_option_texts, write_xlsx_node

    (tmp_path / "inputs").mkdir()
    (tmp_path / "outputs").mkdir()

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["C10"] = "1. 甲选项全文"
    ws["C11"] = "2. 乙选项全文"
    wb.save(tmp_path / "inputs" / "collection_template.xlsx")
    wb.close()

    catalog = {
        "sheets": [
            {
                "name": "Sheet1",
                "sections": [
                    {
                        "indicators": [
                            {
                                "indicator_name": "示例组",
                                "choice_mode": "exclusive",
                                "rows": [
                                    {"row": 10, "option_text": "1. 甲选项全文"},
                                    {"row": 11, "option_text": "2. 乙选项全文"},
                                ],
                            }
                        ]
                    }
                ],
            }
        ]
    }
    (tmp_path / "outputs" / "template_row_catalog.json").write_text(
        json.dumps(catalog, ensure_ascii=False), encoding="utf-8"
    )

    assert _exclusive_option_texts(tmp_path) == {
        ("Sheet1", 10): "1. 甲选项全文",
        ("Sheet1", 11): "2. 乙选项全文",
    }

    # 模型只说「选中」，且故意给了个被截断的错文本，都应被模板原文覆盖
    write_xlsx_node(
        {
            "extract_root": str(tmp_path),
            "row_fills": [
                {"sheet": "Sheet1", "row": 11, "choice": "选中", "remark": "理由"},
            ],
        }
    )

    check = load_workbook(tmp_path / "outputs" / "collection_filled.xlsx")
    try:
        assert check["Sheet1"]["D11"].value == "2. 乙选项全文"
        assert check["Sheet1"]["D10"].value is None
    finally:
        check.close()


def test_case1_fill_group_has_an_output_cap():
    """case1 填表曾完全没有输出上限，模型失控生成会拖垮整个任务。"""
    from app.agents.nodes.case1 import fill_group

    assert fill_group._FILL_GROUP_MAX_TOKENS >= 8192


def test_ocr_meta_fingerprint_invalidates_when_code_or_switches_change(tmp_path: Path):
    """OCR 代码或开关变了，历史 ocr_text 不能再被复用。

    复用只按材料 md5 匹配，不看产物是哪版代码跑的；改完 OCR 链路后同材料的
    新任务会静默复用旧 markdown，日志只有一句 OCR skipped，产物上看不出来。
    """
    import ocr_meta

    ocr_dir = tmp_path / "ocr_text"
    switches = {"text_layer_guard": True, "table_split": True}

    assert ocr_meta.meta_matches(ocr_dir, switches) is False  # 没有 meta 一律不复用

    ocr_meta.write_meta(ocr_dir, switches)
    assert ocr_meta.meta_matches(ocr_dir, switches) is True

    # 开关变了就失效
    assert ocr_meta.meta_matches(ocr_dir, {"text_layer_guard": False, "table_split": True}) is False

    # 代码变了就失效（直接改存档里的指纹等价于改代码）
    import json

    meta_path = ocr_dir / ocr_meta.META_NAME
    data = json.loads(meta_path.read_text(encoding="utf-8"))
    data["code"] = "0" * 16
    meta_path.write_text(json.dumps(data), encoding="utf-8")
    assert ocr_meta.meta_matches(ocr_dir, switches) is False


def test_ocr_code_fingerprint_is_stable_and_covers_ocr_scripts():
    import ocr_meta

    assert ocr_meta.code_fingerprint() == ocr_meta.code_fingerprint()
    assert "table_split.py" in ocr_meta._FINGERPRINT_SOURCES
    assert "text_layer_guard.py" in ocr_meta._FINGERPRINT_SOURCES


def test_calc_sum_partial_equals_existing_value_applies_rule_for_provenance():
    """部分和与模型已填值相等时应用规则，只补溯源不改数。

    2026-07-31 任务 90d60117：源报表没有「应付票据」这一行，模型把「应付账款」
    的值同时填进合并行「应付票据及应付账款」，但只给应付账款带了 evidence_refs。
    计算规则因「来源不全」跳过、不打标记，合并行于是既无证据又无计算规则标记，
    被硬门禁「有值但没有来源证据」判定整个任务失败。

    空 ≠ 缺失：那一行在源报表里本就不存在，部分和就是全和。
    值不同时仍保留模型值（见 test_calc_sum_partial_sources_fill_empty_but_never_overwrite）。
    """
    calc = _calc_module()
    data = {
        "sheets": [
            {
                "sheet": "资产负债表",
                "items": [
                    {
                        "item_id": "r64",
                        "label": "应付票据及应付账款",
                        "evidence_refs": [],
                        "reason_one_line": "文件中未发现相关信息",
                        "fields": {"C": {"cell": "C64", "value": 1758823768.78}},
                    },
                    {
                        "item_id": "r65",
                        "label": "其中：应付票据",
                        "fields": {"C": {"cell": "C65", "value": None}},
                    },
                    {
                        "item_id": "r66",
                        "label": "应付账款",
                        "evidence_refs": ["ocr_text/sources/新鸿集团审计报告2023.pdf.md"],
                        "fields": {"C": {"cell": "C66", "value": 1758823768.78}},
                    },
                ],
            }
        ]
    }
    report = calc.apply_calc_rules(
        data,
        [
            {
                "op": "sum",
                "sheet": "资产负债表",
                "target_label": "应付票据及应付账款",
                "source_labels": ["其中：应付票据", "应付账款"],
            }
        ],
    )

    target = data["sheets"][0]["items"][0]
    assert target["fields"]["C"]["value"] == 1758823768.78        # 值不变
    assert target["reason_one_line"].startswith("计算规则(")        # 拿到标记
    assert target["evidence_refs"] == [                            # 继承来源证据
        "ocr_text/sources/新鸿集团审计报告2023.pdf.md"
    ]
    # 值本来就相等，不算「覆盖模型值」，不该出那条复核提示
    assert not any(
        f["kind"] == "calc_overwrote_model_value" for f in report["review_flags"]
    )


def test_carryforward_check_tolerates_both_report_date_conventions():
    """年初列的 report_date 口径不统一，两种配对都要试，否则会把正确数据判成颠倒。

    2026-07-31 任务 dcb2e2bd：同一份 2023 报告里，年初列的事实一部分标成
    「该列所属期」（2022-12-31），一部分标成「报表日」（2023-12-31）。原实现只
    按「上年期末」比，把前者判成 8/8 全不一致，报出「期末/年初两列可能被识别颠倒」
    ——而那批数值经逐格核对全部正确。
    """
    from app.agents.nodes.case2.evidence import _carryforward_review_flags

    def fact(date_str, period, subject, value):
        return {
            "statement_name": "资产负债表",
            "report_date": date_str,
            "source_period": period,
            "subject_name": subject,
            "value": value,
        }

    subjects = {"货币资金": 800955434.17, "短期借款": 767700000.00, "应付账款": 730625997.92}
    facts = []
    # 口径一：年初列标成该列所属期 -> 应与「同日期末」一致
    for name, value in subjects.items():
        facts.append(fact("2022-12-31", "年初余额", name, value))
        facts.append(fact("2022-12-31", "期末余额", name, value))
    # 口径二：年初列标成报表日 -> 应与「上年期末」一致
    for name, value in subjects.items():
        facts.append(fact("2023-12-31", "年初余额", name, value))

    assert _carryforward_review_flags({"facts": facts}) == []


def test_carryforward_check_still_catches_a_real_swap():
    """两种配对都对不上时仍要报警，否则这道体检就白设了。"""
    from app.agents.nodes.case2.evidence import _carryforward_review_flags

    def fact(date_str, period, subject, value, src):
        return {
            "statement_name": "资产负债表",
            "report_date": date_str,
            "source_period": period,
            "subject_name": subject,
            "value": value,
            "source_ref": src,
        }

    facts = []
    for name, opening, closing in (
        ("货币资金", 111.0, 999.0),
        ("短期借款", 222.0, 888.0),
        ("应付账款", 333.0, 777.0),
        ("存货", 444.0, 666.0),
    ):
        # 2022 年报的年初列（疑似与期末颠倒），拿另一份 2021 年报的期末数交叉
        facts.append(fact("2022-12-31", "年初余额", name, opening, "2022.md"))
        facts.append(fact("2022-12-31", "期末余额", name, closing, "2022.md"))
        facts.append(fact("2021-12-31", "期末余额", name, closing * 2, "2021.md"))

    flags = _carryforward_review_flags({"facts": facts})
    assert [f["kind"] for f in flags] == ["carryforward_mismatch"]


def _grounding_rows(tmp_path: Path, name: str, text: str):
    from app.agents.nodes.case2.evidence import _source_line_index

    (tmp_path / name).write_text(text, encoding="utf-8")
    return _source_line_index(tmp_path, name)


def test_fact_grounding_rejects_value_borrowed_from_adjacent_row(tmp_path: Path):
    """科目名与数值必须在源文档同一行，否则这条事实没有落地依据。

    2026-07-31 任务 dcb2e2bd：主表第 145 行「递延所得税资产」三格全空，
    第 146 行「其他非流动资产」= 18,763,198.09。模型把两行拼在一起，还编了一条
    evidence_text «递延所得税资产 18,763,198.09»——该字符串在原文中并不存在，
    只核对模型自报的证据拦不住，必须回源文档核。
    """
    from app.agents.nodes.case2.evidence import _fact_is_grounded

    rows = _grounding_rows(
        tmp_path,
        "bs.md",
        "## 合并资产负债表\n"
        "| 长期待摊费用 | 872,208.71 | 868,635.47 |\n"
        "| 递延所得税资产 |  |  |\n"
        "| 其他非流动资产 | 18,763,198.09 | 1,440,060.92 |\n",
    )

    borrowed = {"subject_name": "递延所得税资产", "value": 18763198.09}
    genuine = {"subject_name": "其他非流动资产", "value": 18763198.09}

    assert _fact_is_grounded(borrowed, rows) is False
    assert _fact_is_grounded(genuine, rows) is True


def test_fact_grounding_tolerates_label_and_separator_differences(tmp_path: Path):
    """模板标签与源文档写法不同、OCR 分隔符读错，都不能算成无据。"""
    from app.agents.nodes.case2.evidence import _fact_is_grounded

    rows = _grounding_rows(
        tmp_path,
        "bs.md",
        "## 合并资产负债表\n"
        "| 固定资产清理 |  |  | 实收资本 | 100,000,000.00 | 100,000,000.00 |\n"
        "| 货币资金 | 399,093,581.74 | 800.955.434.17 |\n"
        "| 其中：应付票据 | 1,234.00 |  |\n",
    )

    # 模板标签带括号后缀，源文档没有
    assert _fact_is_grounded({"subject_name": "实收资本(或股本)", "value": 100000000.0}, rows) is True
    # 千分位被 OCR 读成句点
    assert _fact_is_grounded({"subject_name": "货币资金", "value": 800955434.17}, rows) is True
    # 「其中：」前缀
    assert _fact_is_grounded({"subject_name": "应付票据", "value": 1234.00}, rows) is True


def test_fact_grounding_is_lenient_when_source_unavailable():
    """取不到源文档时不做判断，不冤枉。"""
    from app.agents.nodes.case2.evidence import _fact_is_grounded

    assert _fact_is_grounded({"subject_name": "货币资金", "value": 1.0}, []) is True


def test_facts_for_fill_excludes_ungrounded_facts():
    from app.agents.nodes.case2.evidence import facts_for_fill

    catalog = {
        "facts": [
            {"item_id": "资产负债表:r6", "subject_name": "货币资金", "value": 1.0},
            {
                "item_id": "资产负债表:r51",
                "subject_name": "递延所得税资产",
                "value": 2.0,
                "grounded": False,
            },
        ]
    }
    got = facts_for_fill(
        catalog,
        {"资产负债表:r6", "资产负债表:r51"},
        sheet_name="资产负债表",
        period_mapping={},
    )
    assert [f["item_id"] for f in got] == ["资产负债表:r6"]


def test_fact_grounding_rejects_notes_value_when_main_statement_lists_the_subject(tmp_path: Path):
    """附注明细表里如实抄来的行，也不能当主表科目的依据。

    2026-07-31 任务 dcb2e2bd：附注「（九）其他非流动资产」的构成里有一行
    「递延所得税资产 | 18,763,198.09」，模型照抄给了主表科目 r51，导致
    非流动资产合计重复计入。而主表第 145 行确实列了「递延所得税资产」，只是为空。
    """
    from app.agents.nodes.case2.evidence import _fact_is_grounded

    rows = _grounding_rows(
        tmp_path,
        "bs.md",
        "## 合并资产负债表\n"
        "| 递延所得税资产 |  |  |\n"
        "| 其他非流动资产 | 18,763,198.09 | 1,440,060.92 |\n"
        "## （九）其他非流动资产\n"
        "| 项目 | 年末余额 | 年初余额 |\n"
        "| 递延所得税资产 | 18,763,198.09 | 1,440,060.92 |\n",
    )

    # 主表列了该科目（虽为空）→ 附注的值不得采用
    assert _fact_is_grounded({"subject_name": "递延所得税资产", "value": 18763198.09}, rows) is False
    # 主表自己那一行的值照常放行
    assert _fact_is_grounded({"subject_name": "其他非流动资产", "value": 18763198.09}, rows) is True


def test_fact_grounding_allows_notes_when_main_statement_omits_the_subject(tmp_path: Path):
    """主表压根没有该科目时，附注可以兜底——有些科目只在附注披露，一律拒绝会误伤。"""
    from app.agents.nodes.case2.evidence import _fact_is_grounded

    rows = _grounding_rows(
        tmp_path,
        "bs.md",
        "## 合并资产负债表\n"
        "| 货币资金 | 399,093,581.74 |  |\n"
        "## （十二）其他说明\n"
        "| 受限资金 | 12,345.67 |  |\n",
    )

    assert _fact_is_grounded({"subject_name": "受限资金", "value": 12345.67}, rows) is True


_SHIFTED_TABLE = (
    # 表头被并进首个科目行（连同骑缝章号），其后每行的科目名都比数值晚一行
    "## 利润表\n"
    "|一、营业收入 230382||本年累计金额|本月金额|\n"
    "| 减：营业成本 | 1 | 100.00 | 10.00 |\n"
    "| 税金及附加 | 2 | 200.00 | 20.00 |\n"
    "| 销售费用 | 3 | 300.00 | 30.00 |\n"
    "| 管理费用 | 4 | 400.00 | 40.00 |\n"
    "| 财务费用 | 6 | 500.00 | 50.00 |\n"
    "| 研发费用 | 7 | 600.00 | 60.00 |\n"
)
_ALIGNED_TABLE = (
    "## 仅限用于XX资产尽调 利润表\n"
    "| 一、营业收入 | 1 | 11.00 | 1.00 |\n"
    "| 减：营业成本 | 2 | 22.00 | 2.00 |\n"
)
_SHIFTED_FACTS = [
    {"subject_name": "一、营业收入", "value": 100.0},
    {"subject_name": "减：营业成本", "value": 200.0},
    {"subject_name": "税金及附加", "value": 300.0},
    {"subject_name": "销售费用", "value": 400.0},
    {"subject_name": "管理费用", "value": 500.0},
    {"subject_name": "财务费用", "value": 600.0},
]


def test_row_offset_detected_per_statement_not_per_file(tmp_path: Path):
    """一份材料里错行的表和对齐的表并存时，偏移必须按单张报表统计。

    东厦四个报告期各一张利润表，只有 2020 那张的表头被并进首行导致整表错行。
    按整个文件统计会被三张对齐的表投票压过去，那张表的事实全部误杀。
    """
    from app.agents.nodes.case2.evidence import _detect_row_offsets, _fact_is_grounded

    rows = _grounding_rows(tmp_path, "pl.md", _SHIFTED_TABLE + _ALIGNED_TABLE)
    facts = _SHIFTED_FACTS + [{"subject_name": "一、营业收入", "value": 11.0}]
    offsets = _detect_row_offsets(facts, rows)

    assert all(_fact_is_grounded(f, rows, row_offsets=offsets) for f in _SHIFTED_FACTS)
    # 对齐的那张表不受影响，仍按同行判定
    assert _fact_is_grounded(facts[-1], rows, row_offsets=offsets) is True


def test_row_offset_does_not_rescue_a_single_borrowed_value(tmp_path: Path):
    """单条事实错配形不成整表规律，仍应判无据——否则 dcb2e2bd 的编造会被放回来。"""
    from app.agents.nodes.case2.evidence import _detect_row_offsets, _fact_is_grounded

    rows = _grounding_rows(
        tmp_path,
        "bs.md",
        "## 合并资产负债表\n"
        "| 长期待摊费用 | 872,208.71 | 868,635.47 |\n"
        "| 递延所得税资产 |  |  |\n"
        "| 其他非流动资产 | 18,763,198.09 | 1,440,060.92 |\n",
    )
    borrowed = {"subject_name": "递延所得税资产", "value": 18763198.09}
    offsets = _detect_row_offsets([borrowed], rows)

    assert _fact_is_grounded(borrowed, rows, row_offsets=offsets) is False


def test_statement_head_survives_watermark_text(tmp_path: Path):
    """水印、骑缝章文字被 OCR 并进标题行时，该表仍算主表区。

    东厦 2021 页读成「## 仅限用于渐商资产东清(-d2o地块能资内准 利润表」，
    按「标题等于报表名」匹配会让整张表落在主表区外，该页事实全部判无据。
    """
    from app.agents.nodes.case2.evidence import _source_line_index

    rows = _source_line_index(*_write(tmp_path, "pl.md", _ALIGNED_TABLE))
    assert any(row[2] for row in rows)
    # 附注类标题仍然不算主表区
    notes = _source_line_index(*_write(tmp_path, "n.md", "## 利润表附注明细\n| 甲 | 1.00 |\n"))
    assert not any(row[2] for row in notes)


def _write(tmp_path: Path, name: str, text: str):
    (tmp_path / name).write_text(text, encoding="utf-8")
    return tmp_path, name


def test_carryforward_check_does_not_mix_sources_in_one_bucket():
    """同一期次的年初桶会收到多份报告的事实，各报告口径不同，混在一起比谁都对不上。

    2026-07-31 任务 a9bb0249：「2022-12-31 年初」桶里，货币资金来自 2023 报告
    （该列所属期口径，实为 2022 年末），应付账款来自 2022 报告（报表日口径，
    实为 2021 年末），8/14 对不上。按来源文档分桶后两边各自都能对上。
    """
    from app.agents.nodes.case2.evidence import _carryforward_review_flags

    def fact(date_str, period, subject, value, src):
        return {
            "statement_name": "资产负债表",
            "report_date": date_str,
            "source_period": period,
            "subject_name": subject,
            "value": value,
            "source_ref": src,
        }

    y2022 = {"货币资金": 800955434.17, "存货": 10212971620.60, "短期借款": 767700000.00}
    y2021 = {"货币资金": 1094162068.88, "存货": 11527112451.13, "短期借款": 829800000.00}

    facts = []
    for name, value in y2022.items():
        facts.append(fact("2022-12-31", "期末余额", name, value, "2022.md"))
        # 2023 报告的年初列：该列所属期口径 -> 与同日期末一致
        facts.append(fact("2022-12-31", "年初余额", name, value, "2023.md"))
    for name, value in y2021.items():
        facts.append(fact("2021-12-31", "期末余额", name, value, "2021.md"))
        # 2022 报告的年初列：报表日口径 -> 与上年期末一致
        facts.append(fact("2022-12-31", "年初余额", name, value, "2022.md"))

    assert _carryforward_review_flags({"facts": facts}) == []


def test_chunk_routing_matches_labels_with_parenthetical_suffix():
    """模板标签带括号后缀、源文档没有时，该指标仍要被派给含它的分块。

    2026-07-31 任务 a9bb0249：模板「实收资本(或股本)」归一化后是「实收资本或股本」，
    源文档写的是「实收资本」，路由的 _label_variants 只做了「去其中/加/减前缀」、
    没做「去括号后缀」，于是 2021 报告的主表块只分到 67 个指标（其余两份 121 个），
    r106 不在其中，那一期整个漏抽。路由与落地校验现共用 _subject_keys。
    """
    from app.agents.nodes.case2.evidence import _route_evidence_chunks, _subject_keys

    assert "实收资本" in _subject_keys("实收资本(或股本)")
    assert "应付票据" in _subject_keys("其中：应付票据")

    chunks = [
        {
            "source_ref": "ocr_text/sources/bs.md",
            "text": "## 合并资产负债表\n| 固定资产清理 |  |  | 实收资本 | 100,000,000.00 |\n",
        }
    ]
    targets = [
        {"item_id": "资产负债表:r106", "label": "实收资本(或股本)", "sheet_name": "资产负债表"},
        {"item_id": "资产负债表:r6", "label": "货币资金", "sheet_name": "资产负债表"},
    ]
    routed, _ = _route_evidence_chunks(chunks, targets)

    ids = {str(t.get("item_id")) for t in (routed[0].get("targets") or [])}
    assert "资产负债表:r106" in ids


def test_carryforward_check_never_compares_one_report_against_itself():
    """同一份报告的年初列与期末列本就是两期，不能互比。

    2026-07-31 任务 08adc5c3：材料只有三年，2020 年报不存在，「2021-12-31 年初」
    唯一的候选对手就是同一份 2021 报告的期末列，于是报出 26/29 不一致。
    """
    from app.agents.nodes.case2.evidence import _carryforward_review_flags

    def fact(period, subject, value):
        return {
            "statement_name": "资产负债表",
            "report_date": "2021-12-31",
            "source_period": period,
            "subject_name": subject,
            "value": value,
            "source_ref": "2021.md",
        }

    facts = []
    for name, opening, closing in (
        ("货币资金", 1575470029.83, 1094162068.88),
        ("短期借款", 250000000.00, 829800000.00),
        ("应付账款", 111.0, 222.0),
        ("存货", 333.0, 444.0),
    ):
        facts.append(fact("年初余额", name, opening))
        facts.append(fact("期末余额", name, closing))

    assert _carryforward_review_flags({"facts": facts}) == []
